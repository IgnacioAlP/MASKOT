import os
import sys
import json
import time
import logging
import hashlib
import random
import re
from datetime import datetime, date, timezone, timedelta
from functools import wraps
import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter    
import traceback

from flask import (
    Flask, render_template, request, redirect, url_for, 
    session, flash, jsonify, send_from_directory, send_file
)
from werkzeug.utils import secure_filename

# Conexión a Base de Datos
from bd import obtener_conexion, obtener_tenant_id

# Importar controladores
from controladores import (
    usuarios_controlador,
    citas_controlador,
    servicios_controlador,
    personal_controlador,
    productos_controlador,
    asistencia_controlador,
    clientes_controlador,
    compras_controlador,
    ventas_controlador,
    mascotas_controlador
)
from controladores import fidelizacion_controlador as fidelizacion_ctrl

# ─── CONFIGURACIÓN DE LA APLICACIÓN ──────────────────────────────────────────

BASE_DIR = os.path.dirname(os.path.abspath(__file__))  # MASKOT/api
ROOT_DIR = os.path.dirname(BASE_DIR)                    # MASKOT
ZONA_HORARIA_PERU = timezone(timedelta(hours=-5))

app = Flask(
    __name__,
    template_folder=os.path.join(ROOT_DIR, 'templates'),
    static_folder=os.path.join(ROOT_DIR, 'static'),
    static_url_path='/static'
)

app.secret_key = os.environ.get('SECRET_KEY', 'super_secret_key_veterinaria')
use_secure_cookies = os.environ.get('FLASK_ENV', '').lower() == 'production' or os.environ.get('USE_SECURE_COOKIES', '') == 'True'
app.config['SESSION_COOKIE_SECURE'] = bool(use_secure_cookies)
app.config['SESSION_COOKIE_HTTPONLY'] = True
app.config['SESSION_COOKIE_SAMESITE'] = 'Lax'

logging.basicConfig(level=logging.INFO, format='%(asctime)s [%(levelname)s] %(message)s')
logger = logging.getLogger(__name__)

ALLOWED_EXTENSIONS = {'png', 'jpg', 'jpeg', 'gif', 'webp', 'pdf'}

if os.environ.get('VERCEL'):
    UPLOAD_FOLDER = '/tmp/uploads'
else:
    UPLOAD_FOLDER = os.path.join(os.path.dirname(__file__), '..', 'static', 'images')

app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

try:
    os.makedirs(UPLOAD_FOLDER, exist_ok=True)
except Exception:
    pass

RECENT_SCANS = []


def _ensure_schema():
    """Migración liviana para ajustar el esquema de Supabase/PostgreSQL."""
    column_migrations = [
        ("productos", "codigo_barra", "ALTER TABLE productos ADD COLUMN IF NOT EXISTS codigo_barra VARCHAR(100) DEFAULT NULL"),
        ("productos", "stock_minimo", "ALTER TABLE productos ADD COLUMN IF NOT EXISTS stock_minimo INT DEFAULT 5"),
        ("citas",     "tenant_id",   "ALTER TABLE citas ADD COLUMN IF NOT EXISTS tenant_id INT DEFAULT 1"),
        ("clientes",  "documento",   "ALTER TABLE clientes ADD COLUMN IF NOT EXISTS documento VARCHAR(50) DEFAULT NULL"),
        ("clientes",  "activo",      "ALTER TABLE clientes ADD COLUMN IF NOT EXISTS activo BOOLEAN DEFAULT true"),
        ("ventas",    "tenant_id",   "ALTER TABLE ventas ADD COLUMN IF NOT EXISTS tenant_id INT DEFAULT 1"),
        ("ventas",    "cliente_id",  "ALTER TABLE ventas ADD COLUMN IF NOT EXISTS cliente_id INT DEFAULT NULL"),
        ("ventas",    "cliente_nombre", "ALTER TABLE ventas ADD COLUMN IF NOT EXISTS cliente_nombre VARCHAR(255) DEFAULT 'Cliente General'"),
        ("ventas",    "cliente_documento", "ALTER TABLE ventas ADD COLUMN IF NOT EXISTS cliente_documento VARCHAR(50) DEFAULT NULL"),
        ("ventas",    "vendedor_id", "ALTER TABLE ventas ADD COLUMN IF NOT EXISTS vendedor_id INT DEFAULT NULL"),
        ("ventas",    "vendedor_nombre", "ALTER TABLE ventas ADD COLUMN IF NOT EXISTS vendedor_nombre VARCHAR(100) DEFAULT NULL"),
        ("ventas",    "productos",   "ALTER TABLE ventas ADD COLUMN IF NOT EXISTS productos TEXT DEFAULT '[]'"),
    ]
    try:
        conn = obtener_conexion()
        with conn.cursor() as cursor:
            for table, column, sql in column_migrations:
                cursor.execute("""
                    SELECT 1 FROM information_schema.columns 
                    WHERE table_name = %s AND column_name = %s
                """, (table, column))
                if not cursor.fetchone():
                    try:
                        cursor.execute(sql)
                        conn.commit()
                        logger.info(f"Schema migration: columna '{column}' añadida a {table}.")
                    except Exception as ex_col:
                        conn.rollback()
                        logger.warning(f"No se pudo ejecutar {sql}: {ex_col}")

            try:
                cursor.execute("ALTER TABLE clientes ALTER COLUMN email DROP NOT NULL;")
                conn.commit()
            except Exception:
                conn.rollback()

            for col in ['vendedor_nombre', 'vendedor_id', 'cliente_id', 'cliente_nombre']:
                try:
                    cursor.execute("""
                        SELECT is_nullable FROM information_schema.columns 
                        WHERE table_name = 'ventas' AND column_name = %s
                    """, (col,))
                    v_col = cursor.fetchone()
                    if v_col and (v_col[0] or '').upper() == 'NO':
                        cursor.execute(f"ALTER TABLE ventas ALTER COLUMN {col} DROP NOT NULL")
                        conn.commit()
                except Exception:
                    conn.rollback()

        conn.close()
    except Exception as e:
        logger.warning(f"Schema migration check failed: {e}")


try:
    _ensure_schema()
except Exception:
    pass


@app.route('/health', methods=['GET'])
def health_check():
    """Endpoint liviano para confirmar que Vercel puede cargar la aplicación."""
    return jsonify({
        'status': 'ok',
        'service': 'maskot-api',
        'environment': 'vercel' if os.environ.get('VERCEL') else 'local'
    }), 200


@app.errorhandler(Exception)
def handle_exception(e):
    logger.error(f"Error no controlado: {e}")
    return render_template('error.html'), 500


# ─── FILTROS Y CONTEXT PROCESSORS ─────────────────────────────────────────────

@app.template_filter('dateformat')
def dateformat(value, format='%d/%m/%Y'):
    if value is None:
        return ""
    if isinstance(value, str):
        try:
            value = datetime.strptime(value, '%Y-%m-%d').date()
        except Exception:
            return value
    elif isinstance(value, datetime):
        value = value.date()
    
    if hasattr(value, 'strftime'):
        return value.strftime(format)
    return str(value)


@app.template_filter('days_until')
def days_until(value):
    if value is None:
        return 999
    if isinstance(value, str):
        try:
            value = datetime.strptime(value, '%Y-%m-%d').date()
        except Exception:
            return 999
    elif isinstance(value, datetime):
        value = value.date()
    
    if hasattr(value, '__sub__'):
        today_d = datetime.now(ZONA_HORARIA_PERU).date()
        delta = value - today_d
        return delta.days
    return 999


@app.context_processor
def utility_processor():
    def format_date(date_value, format='%d/%m/%Y'):
        if date_value is None:
            return ""
        if isinstance(date_value, str):
            try:
                date_value = datetime.strptime(date_value, '%Y-%m-%d').date()
            except Exception:
                return date_value
        elif isinstance(date_value, datetime):
            date_value = date_value.date()
        
        if hasattr(date_value, 'strftime'):
            return date_value.strftime(format)
        return str(date_value)
    
    def today():
        return datetime.now(ZONA_HORARIA_PERU).date()
    
    def today_string():
        return datetime.now(ZONA_HORARIA_PERU).strftime('%Y-%m-%d')
    
    def days_difference(date1, date2=None):
        if date2 is None:
            date2 = datetime.now(ZONA_HORARIA_PERU).date()
        if isinstance(date1, str):
            try:
                date1 = datetime.strptime(date1, '%Y-%m-%d').date()
            except Exception:
                return 999
        elif isinstance(date1, datetime):
            date1 = date1.date()
        return (date1 - date2).days
    
    def calculate_work_hours(entrada, salida):
        if not entrada or not salida or str(salida).lower() in ['en turno', '--:--']:
            return 0
        try:
            if isinstance(entrada, str):
                entrada = datetime.strptime(entrada.strip(), '%H:%M:%S').time() if ':' in entrada else datetime.strptime(entrada.strip(), '%I:%M %p').time()
            if isinstance(salida, str):
                salida = datetime.strptime(salida.strip(), '%H:%M:%S').time() if ':' in salida else datetime.strptime(salida.strip(), '%I:%M %p').time()
            
            today_date = datetime.now(ZONA_HORARIA_PERU).date()
            entrada_dt = datetime.combine(today_date, entrada)
            salida_dt = datetime.combine(today_date, salida)
            
            delta = salida_dt - entrada_dt
            return max(0.0, round(delta.total_seconds() / 3600, 1))
        except Exception:
            return 0
    
    return dict(
        format_date=format_date, 
        today=today, 
        today_string=today_string, 
        days_difference=days_difference,
        calculate_work_hours=calculate_work_hours
    )


@app.context_processor
def cart_context():
    cart = session.get('cart', {})
    total_qty = 0
    try:
        for v in cart.values():
            total_qty += int(v.get('qty', 0))
    except Exception:
        total_qty = 0
    return dict(cart_count=total_qty)


@app.context_processor
def inject_global_datetime():
    ahora = datetime.now(ZONA_HORARIA_PERU)
    return {
        'momento_actual': ahora,
        'today': lambda: ahora.date(),
        'now': ahora,
        'format_date': lambda d: d.strftime('%d/%m/%Y') if hasattr(d, 'strftime') else str(d)
    }


def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


def requiere_autenticacion(roles_permitidos=None):
    def decorator(f):
        @wraps(f)
        def wrapper(*args, **kwargs):
            if 'usuario' not in session or 'rol' not in session:
                flash('Debes iniciar sesión para acceder a esta página.', 'warning')
                return redirect(url_for('login'))
            
            if session['rol'] == 'superadmin':
                return f(*args, **kwargs)
            
            if roles_permitidos and session['rol'] not in roles_permitidos:
                flash('No tienes permisos para acceder a esta página.', 'error')
                return redirect(url_for('dashboard'))
            
            return f(*args, **kwargs)
        return wrapper
    return decorator


@app.route('/favicon.ico')
def favicon():
    return send_from_directory(os.path.join(app.root_path, 'static', 'img'),
                               'maskot-favicon.ico', mimetype='image/vnd.microsoft.icon')


@app.route('/')
def index():
    if 'usuario' in session:
        return redirect(url_for('dashboard'))
    return redirect(url_for('index.html'))


# ─── AUTENTICACIÓN Y SESIÓN ──────────────────────────────────────────────────

@app.route('/login', methods=['GET', 'POST'])
def login():
    if 'rol' in session and session.get('rol') in ['dueño', 'admin', 'empleado', 'superadmin']:
        return redirect(url_for('dashboard'))

    if request.method == 'POST':
        username = (request.form.get('username') or '').strip()
        password = (request.form.get('password') or '').strip()
        next_url = request.args.get('next') or request.form.get('next')

        if not username or not password:
            flash('Por favor complete todos los campos.', 'warning')
            return render_template('login.html')
        
        try:
            auth = usuarios_controlador.verificar_credenciales(username, password)
        except Exception as e:
            logger.error(f"Error en autenticación: {e}")
            flash('Error interno del sistema', 'error')
            return render_template('login.html')
            
        if not auth.get('success'):
            flash(auth.get('message', 'Credenciales inválidas'), 'error')
            return render_template('login.html')
            
        user = auth.get('user')
        session.permanent = True
        session['user_id'] = user[0]
        session['usuario_id'] = user[0]
        session['usuario'] = user[1]
        session['rol'] = user[3]
        session['tenant_id'] = user[5] if len(user) > 5 and user[5] else 1
        
        flash(f"¡Bienvenido de nuevo, {user[1]}!", 'success')
        if next_url and next_url.startswith('/'):
            return redirect(next_url)
        
        if user[3] == 'superadmin':
            return redirect(url_for('gestionar_tenants'))
            
        return redirect(url_for('dashboard'))
            
    return render_template('login.html')


@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('index'))


# ─── DASHBOARD ───────────────────────────────────────────────────────────────

def obtener_fecha_hoy_peru():
    return datetime.now(ZONA_HORARIA_PERU).strftime('%Y-%m-%d')

@app.context_processor
def inject_today():
    return dict(today_string=obtener_fecha_hoy_peru)


def obtener_columna_fecha(cursor):
    cursor.execute("""
        SELECT column_name 
        FROM information_schema.columns 
        WHERE table_name = 'ventas' 
          AND column_name IN ('fecha_venta', 'fecha', 'created_at', 'fecha_registro', 'fechahora')
        ORDER BY CASE column_name
            WHEN 'fecha_venta' THEN 1
            WHEN 'fecha' THEN 2
            WHEN 'created_at' THEN 3
            ELSE 4
        END
        LIMIT 1;
    """)
    res = cursor.fetchone()
    return res[0] if res else 'fecha_venta'


@app.route('/dashboard')
def dashboard():
    if 'rol' not in session:
        return redirect(url_for('login'))
    
    rol = session['rol']
    hoy = obtener_fecha_hoy_peru()
    
    citas_hoy = []
    alertas_fidelizacion = []
    
    cuadre_hoy = {
        'cantidad_ventas': 0,
        'efectivo': 0.0,
        'yape': 0.0,
        'tarjeta': 0.0,
        'total': 0.0
    }
    totales_dia = cuadre_hoy.copy()
    
    if rol in ['admin', 'empleado', 'dueño']:
        try:
            citas_hoy = citas_controlador.obtener_citas_por_fecha(hoy)
        except Exception as e:
            logger.warning(f"Error obteniendo citas de hoy: {e}")
            
        try:
            if hasattr(fidelizacion_ctrl, 'sincronizar_fidelizacion_desde_citas'):
                fidelizacion_ctrl.sincronizar_fidelizacion_desde_citas()
            if hasattr(fidelizacion_ctrl, 'obtener_alertas_recientes'):
                alertas_fidelizacion = fidelizacion_ctrl.obtener_alertas_recientes(limit=10)
        except Exception as e:
            logger.warning(f"Error fidelización: {e}")

        conexion = None
        try:
            conexion = obtener_conexion()
            with conexion.cursor() as cursor:
                col_fecha = obtener_columna_fecha(cursor)

                cursor.execute(f"""
                    SELECT 
                        COUNT(*) AS cantidad_ventas,
                        COALESCE(SUM(total), 0) AS total_soles,
                        COALESCE(SUM(CASE WHEN LOWER(TRIM(metodo_pago)) = 'efectivo' THEN total ELSE 0 END), 0) AS total_efectivo,
                        COALESCE(SUM(CASE WHEN LOWER(TRIM(metodo_pago)) IN ('yape', 'plin') THEN total ELSE 0 END), 0) AS total_yape,
                        COALESCE(SUM(CASE WHEN LOWER(TRIM(metodo_pago)) NOT IN ('efectivo', 'yape', 'plin') THEN total ELSE 0 END), 0) AS total_tarjeta
                    FROM ventas
                    WHERE DATE({col_fecha}) = %s::date OR {col_fecha}::text LIKE %s || '%%'
                """, (hoy, hoy))
                
                res = cursor.fetchone()
                if res and res[0] is not None:
                    cant = int(res[0])
                    tot_soles = float(res[1])
                    tot_efec = float(res[2])
                    tot_yape = float(res[3])
                    tot_tarj = float(res[4])

                    cuadre_hoy = {
                        'cantidad_ventas': cant,
                        'efectivo': tot_efec,
                        'yape': tot_yape,
                        'tarjeta': tot_tarj,
                        'total': tot_soles
                    }
                    totales_dia = cuadre_hoy.copy()
        except Exception as e:
            if conexion:
                conexion.rollback()
            logger.error(f"Error calculando acumulado de ventas: {e}")
        finally:
            if conexion:
                conexion.close()
    
    if rol == 'dueño':
        total_empleados = 0
        servicios_activos = 0
        productos_bajos = []
        
        try:
            if hasattr(personal_controlador, 'obtener_personal'):
                total_empleados = len(personal_controlador.obtener_personal())
            elif hasattr(personal_controlador, 'obtener_empleados'):
                total_empleados = len(personal_controlador.obtener_empleados())
        except Exception as e:
            logger.warning(f"Error personal dashboard: {e}")

        try:
            if hasattr(servicios_controlador, 'obtener_servicios_activos'):
                servicios_activos = len(servicios_controlador.obtener_servicios_activos())
            elif hasattr(servicios_controlador, 'obtener_servicios'):
                servicios_activos = len(servicios_controlador.obtener_servicios())
        except Exception as e:
            logger.warning(f"Error servicios dashboard: {e}")

        try:
            if hasattr(productos_controlador, 'obtener_productos_bajo_stock'):
                productos_bajos = productos_controlador.obtener_productos_bajo_stock()
        except Exception as e:
            logger.warning(f"Error productos stock bajo dashboard: {e}")
        
        return render_template('dashboard.html', 
                               citas_hoy=citas_hoy, 
                               total_empleados=total_empleados,
                               servicios_activos=servicios_activos, 
                               productos_bajos=productos_bajos,
                               alertas_fidelizacion=alertas_fidelizacion, 
                               totales_dia=totales_dia, 
                               cuadre_hoy=cuadre_hoy)
    
    return render_template('dashboard.html', 
                           citas_hoy=citas_hoy, 
                           alertas_fidelizacion=alertas_fidelizacion, 
                           totales_dia=totales_dia, 
                           cuadre_hoy=cuadre_hoy)


@app.route('/exportar-cuadre-excel', methods=['GET'])
@app.route('/exportar-cierre-diario', methods=['GET'])
def exportar_cierre_diario():
    fecha_filtro = request.args.get('fecha') or obtener_fecha_hoy_peru()
    conexion = None
    try:
        conexion = obtener_conexion()
        with conexion.cursor() as cursor:
            col_fecha = obtener_columna_fecha(cursor)
            
            cursor.execute(f"""
                SELECT id, COALESCE(total, 0.0) AS total, COALESCE(metodo_pago, 'efectivo') AS metodo, {col_fecha} AS fecha
                FROM ventas
                WHERE DATE({col_fecha}::text) = %s::date OR {col_fecha}::text LIKE %s || '%%'
                ORDER BY id ASC
            """, (fecha_filtro, fecha_filtro))

            registros = cursor.fetchall()

        wb = Workbook()
        ws = wb.active
        ws.title = f"Cierre {fecha_filtro}"

        ws.merge_cells("A1:G1")
        ws["A1"] = f"REPORTE DE CIERRE DE CAJA - FECHA: {fecha_filtro}"
        ws["A1"].font = Font(bold=True, size=11)
        ws["A1"].alignment = Alignment(horizontal="center")

        headers = ["N° Venta", "Monto Total", "Efectivo S/", "Yape / Plin S/", "Tarjeta / Otros S/", "Método Pago", "Fecha / Hora"]
        ws.append(headers)

        header_font = Font(bold=True, color="FFFFFF", size=10)
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        border_thin = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

        for col_idx in range(1, 8):
            cell = ws.cell(row=2, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border_thin

        row_start = 3
        t_total, t_efec, t_yape, t_tarj = 0.0, 0.0, 0.0, 0.0

        for idx, row in enumerate(registros, start=1):
            v_id = row[0]
            monto = float(row[1] or 0)
            metodo = str(row[2] or '').strip().lower()
            f_hora = str(row[3] or '')

            efec = monto if metodo == 'efectivo' else 0.0
            yape = monto if metodo in ['yape', 'plin'] else 0.0
            tarj = monto if metodo not in ['efectivo', 'yape', 'plin'] else 0.0

            t_total += monto
            t_efec += efec
            t_yape += yape
            t_tarj += tarj

            ws.append([v_id, monto, efec, yape, tarj, row[2], f_hora])

            for col_idx in range(1, 8):
                ws.cell(row=row_start + idx - 1, column=col_idx).border = border_thin

        ws.append(["TOTALES", t_total, t_efec, t_yape, t_tarj, "", ""])
        last_row = ws.max_row
        
        total_font = Font(bold=True, size=10)
        total_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

        for col_idx in range(1, 8):
            cell = ws.cell(row=last_row, column=col_idx)
            cell.font = total_font
            cell.fill = total_fill
            cell.border = border_thin

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col[1:])
            col_idx = col[0].column
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        nombre_archivo = f"Cierre_Diario_{fecha_filtro}.xlsx"

        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=nombre_archivo
        )

    except Exception as e:
        if conexion:
            conexion.rollback()
        logger.error(f"Error generando reporte de cierre: {e}")
        return f"Error al generar Excel: {str(e)}", 500
    finally:
        if conexion:
            conexion.close()


# ─── VISTA PRINCIPAL DEL PUNTO DE VENTA (POS) ─────────────────────────────────

@app.route('/punto_de_venta', methods=['GET', 'POST'])
@app.route('/pos', methods=['GET', 'POST'])
def punto_de_venta():
    if 'rol' not in session or session['rol'] not in ['admin', 'empleado', 'dueño']:
        flash('Acceso denegado al punto de venta.', 'error')
        return redirect(url_for('dashboard'))
    
    tenant_id = session.get('tenant_id', 1)
    productos_lista = []
    servicios_lista = []

    try:
        conexion = obtener_conexion()
        with conexion.cursor() as cursor:
            cursor.execute("""
                SELECT 
                    id, 
                    nombre, 
                    COALESCE(precio, 0.00) AS precio, 
                    COALESCE(cantidad, 0) AS cantidad, 
                    COALESCE(codigo_barra, '') AS codigo_barra, 
                    'General' AS categoria, 
                    COALESCE(imagen, '') AS imagen 
                FROM productos
                WHERE (tenant_id = %s OR tenant_id IS NULL) 
                  AND COALESCE(activo, true) = true
                  AND LOWER(tipo::text) = 'venta'
                ORDER BY nombre ASC
            """, (tenant_id,))
            
            for r in cursor.fetchall():
                p_id = r[0]
                p_nom = r[1] or ''
                try:
                    p_prec = float(r[2])
                except (ValueError, TypeError):
                    p_prec = 0.0
                try:
                    p_stk = int(r[3])
                except (ValueError, TypeError):
                    p_stk = 0
                
                p_code = str(r[4]) if r[4] else ''
                p_cat = str(r[5])
                p_img = str(r[6]) if r[6] else ''

                productos_lista.append({
                    'id': p_id,
                    'nombre': p_nom,
                    'precio': p_prec,
                    'stock': p_stk,
                    'cantidad': p_stk,
                    'codigo_barra': p_code,
                    'categoria': p_cat,
                    'imagen': p_img,
                    0: p_id, 1: p_nom, 2: p_prec, 3: p_stk, 4: p_code, 5: p_cat, 6: p_img
                })

            cursor.execute("""
                SELECT 
                    id, 
                    nombre, 
                    COALESCE(precio, 0.00) AS precio, 
                    COALESCE(duracion, 0) AS duracion
                FROM servicios
                WHERE (tenant_id = %s OR tenant_id IS NULL) 
                  AND COALESCE(activo, true) = true
                ORDER BY nombre ASC
            """, (tenant_id,))
            
            for s in cursor.fetchall():
                s_id = s[0]
                s_nom = s[1] or ''
                try:
                    s_prec = float(s[2])
                except (ValueError, TypeError):
                    s_prec = 0.0
                s_dur = f"{s[3]} min" if s[3] else ''

                servicios_lista.append({
                    'id': s_id,
                    'nombre': s_nom,
                    'precio': s_prec,
                    'duracion': s_dur,
                    0: s_id, 1: s_nom, 2: s_prec, 3: s_dur, 5: True
                })

        conexion.close()
    except Exception as e:
        logger.error(f"Error consultando productos/servicios para el POS: {e}")

    return render_template('pos.html', productos=productos_lista, servicios=servicios_lista)


# ─── API BÚSQUEDA DE CLIENTES Y PRODUCTOS PARA EL POS ─────────────────────────

@app.route('/api/ventas/buscar-clientes')
def api_buscar_clientes():
    if 'rol' not in session:
        return jsonify({'success': False, 'error': 'No autorizado'}), 401
    
    q = request.args.get('q', '').strip()
    tenant_id = session.get('tenant_id', 1)
    
    try:
        conexion = obtener_conexion()
        with conexion.cursor() as cursor:
            if q:
                param_like = f"%{q}%"
                cursor.execute("""
                    SELECT id, nombre, COALESCE(documento, '') AS documento, COALESCE(telefono, '') AS telefono
                    FROM clientes
                    WHERE (tenant_id = %s OR tenant_id IS NULL)
                      AND COALESCE(activo, true) = true
                      AND (nombre ILIKE %s OR documento ILIKE %s OR telefono ILIKE %s)
                    ORDER BY nombre ASC
                    LIMIT 15
                """, (tenant_id, param_like, param_like, param_like))
            else:
                cursor.execute("""
                    SELECT id, nombre, COALESCE(documento, '') AS documento, COALESCE(telefono, '') AS telefono
                    FROM clientes
                    WHERE (tenant_id = %s OR tenant_id IS NULL)
                      AND COALESCE(activo, true) = true
                    ORDER BY nombre ASC
                    LIMIT 15
                """, (tenant_id,))
            
            rows = cursor.fetchall()
            clientes = [{
                'id': r[0],
                'nombre': r[1],
                'documento': r[2],
                'telefono': r[3]
            } for r in rows]
            
        conexion.close()
        return jsonify({'success': True, 'clientes': clientes})
    except Exception as e:
        logger.error(f"Error buscando clientes para POS: {e}")
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/ventas/buscar-productos')
def api_buscar_productos():
    q = request.args.get('q', '').strip()
    cat_id = request.args.get('categoria_id', '').strip()
    tenant_id = session.get('tenant_id', 1)
    
    try:
        conexion = obtener_conexion()
        with conexion.cursor() as cursor:
            condiciones = [
                "(tenant_id = %s OR tenant_id IS NULL)",
                "COALESCE(activo, true) = true",
                "LOWER(tipo::text) = 'venta'"
            ]
            params = [tenant_id]

            if cat_id and cat_id.isdigit():
                condiciones.append("categoria_id = %s")
                params.append(int(cat_id))

            if q:
                condiciones.append("(nombre ILIKE %s OR codigo_barra::text ILIKE %s OR TRIM(codigo_barra::text) = %s)")
                query_like = f"%{q}%"
                params.extend([query_like, query_like, q])

            where_clause = " AND ".join(condiciones)

            cursor.execute(f"""
                SELECT id, nombre, COALESCE(precio, 0.00) AS precio, COALESCE(cantidad, 0) AS cantidad, 
                       COALESCE(codigo_barra::text, '') AS codigo_barra, 'General' AS categoria
                FROM productos 
                WHERE {where_clause}
                ORDER BY nombre ASC
                LIMIT 50
            """, tuple(params))
            
            rows = cursor.fetchall()
            productos = [{
                'id': r[0],
                'nombre': r[1],
                'precio': float(r[2]) if r[2] is not None else 0.0,
                'stock': int(r[3]) if r[3] is not None else 0,
                'codigo_barra': str(r[4] or ''),
                'categoria': r[5]
            } for r in rows]

        conexion.close()
        return jsonify({'success': True, 'productos': productos})
    except Exception as e:
        logger.error(f"Error buscando productos POS: {e}")
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/ventas/validar-stock', methods=['POST'])
def api_validar_stock():
    data = request.get_json(silent=True) or {}
    items = data.get('items', [])
    errores = []
    
    try:
        conexion = obtener_conexion()
        with conexion.cursor() as cursor:
            for item in items:
                pid = item.get('id')
                cant = int(item.get('cantidad', 1))
                if pid and str(pid).isdigit():
                    cursor.execute("SELECT nombre, cantidad FROM productos WHERE id = %s", (int(pid),))
                    row = cursor.fetchone()
                    if not row:
                        errores.append(f"Producto ID {pid} no encontrado.")
                    elif (row[1] or 0) < cant:
                        errores.append(f"Stock insuficiente para '{row[0]}'. Disponible: {row[1]}, Solicitado: {cant}")
        conexion.close()
        
        if errores:
            return jsonify({'success': False, 'message': 'Validación fallida', 'errors': errores}), 400
        return jsonify({'success': True, 'message': 'Stock verificado correctamente'})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/ventas/procesar', methods=['POST'])
def procesar_venta():
    if 'rol' not in session:
        return jsonify({'success': False, 'error': 'No autorizado'}), 401

    try:
        data = request.get_json(silent=True) or {}
        
        print("\n=================== DEBUG POS ===================")
        print("DATOS RECIBIDOS DEL CLIENTE:", data)
        print("=================================================\n")

        items = data.get('productos') or data.get('items') or []
        if not items:
            return jsonify({'success': False, 'error': 'El carrito está vacío'}), 400

        # 1. Extraer ID del cliente si se envió
        raw_id = data.get('cliente_id') or data.get('id_cliente') or data.get('client_id')
        if isinstance(raw_id, dict):
            raw_id = raw_id.get('id')
        
        cliente_id = None
        if raw_id is not None:
            str_id = str(raw_id).strip()
            if str_id.isdigit() and int(str_id) > 0:
                cliente_id = int(str_id)

        # 2. Extraer Nombre y Documento del cliente
        raw_nombre = (
            data.get('cliente_nombre') or 
            data.get('nombre_cliente') or 
            data.get('client_name') or 
            data.get('cliente') or ''
        )
        if isinstance(raw_nombre, dict):
            raw_doc = raw_nombre.get('documento') or raw_nombre.get('doc') or ''
            raw_nombre = raw_nombre.get('nombre', '')
        else:
            raw_doc = data.get('cliente_documento') or data.get('documento') or data.get('dni') or ''

        cliente_nombre = str(raw_nombre).strip()
        cliente_documento = str(raw_doc).strip() if raw_doc else None

        # Limpiar si trae documento formato "Nombre (DNI: 12345)"
        if '(' in cliente_nombre and ')' in cliente_nombre:
            doc_match = re.search(r'\((.*?)\)', cliente_nombre)
            if doc_match and not cliente_documento:
                cliente_documento = doc_match.group(1).replace('DNI:', '').replace('RUC:', '').strip()
            cliente_nombre = re.sub(r'\s*\(.*?\)', '', cliente_nombre).strip()

        if not cliente_nombre:
            cliente_nombre = 'Cliente General'

        conexion = obtener_conexion()
        tenant_id = session.get('tenant_id', 1)

        with conexion.cursor() as cursor:
            final_cliente_id = None

            # A) Si nos enviaron un ID, verificar su existencia en la tabla clientes
            if cliente_id:
                cursor.execute("""
                    SELECT id, nombre, COALESCE(documento, '') 
                    FROM clientes 
                    WHERE id = %s AND (tenant_id = %s OR tenant_id IS NULL)
                """, (cliente_id, tenant_id))
                cli_db = cursor.fetchone()
                if cli_db:
                    final_cliente_id = cli_db[0]
                    if cli_db[1] and cli_db[1].strip():
                        cliente_nombre = cli_db[1].strip()
                    if cli_db[2] and not cliente_documento:
                        cliente_documento = cli_db[2]

            # B) Si no hay ID pero el cliente no es 'Cliente General', buscar por Nombre/Documento
            if not final_cliente_id and cliente_nombre.lower() != 'cliente general':
                cursor.execute("""
                    SELECT id, nombre, COALESCE(documento, '') 
                    FROM clientes 
                    WHERE LOWER(TRIM(nombre)) = LOWER(TRIM(%s)) 
                      AND (tenant_id = %s OR tenant_id IS NULL) 
                    LIMIT 1
                """, (cliente_nombre, tenant_id))
                cli_db = cursor.fetchone()
                if cli_db:
                    final_cliente_id = cli_db[0]
                    cliente_nombre = cli_db[1]
                    if cli_db[2] and not cliente_documento:
                        cliente_documento = cli_db[2]
                else:
                    # Crear automáticamente el cliente para tener su ID guardado
                    try:
                        email_auto = f"cliente_{int(time.time())}_{random.randint(100,999)}@pos.local"
                        cursor.execute("""
                            INSERT INTO clientes (nombre, email, documento, activo, tenant_id) 
                            VALUES (%s, %s, %s, true, %s) 
                            RETURNING id
                        """, (cliente_nombre, email_auto, cliente_documento, tenant_id))
                        final_cliente_id = cursor.fetchone()[0]
                    except Exception as e_cli:
                        logger.error(f"Error creando cliente automático: {e_cli}")

            # C) Si es 'Cliente General' o falló la asignación previa, vincular con 'Cliente General' en DB si existe
            if not final_cliente_id:
                cursor.execute("""
                    SELECT id FROM clientes 
                    WHERE LOWER(TRIM(nombre)) = 'cliente general' 
                      AND (tenant_id = %s OR tenant_id IS NULL) 
                    LIMIT 1
                """, (tenant_id,))
                cli_gen = cursor.fetchone()
                if cli_gen:
                    final_cliente_id = cli_gen[0]
                else:
                    try:
                        cursor.execute("""
                            INSERT INTO clientes (nombre, email, activo, tenant_id)
                            VALUES ('Cliente General', 'general@pos.local', true, %s)
                            RETURNING id
                        """, (tenant_id,))
                        final_cliente_id = cursor.fetchone()[0]
                    except Exception:
                        final_cliente_id = None

            print(f"-> VENTA FINAL A GUARDAR -> ID Cliente: {final_cliente_id} | Nombre: '{cliente_nombre}' | Doc: '{cliente_documento}'")

            subtotal = float(data.get('subtotal', 0.0))
            igv = float(data.get('igv', 0.0))
            total = float(data.get('total', 0.0))
            monto_recibido = float(data.get('monto_recibido', total))
            cambio = float(data.get('cambio_entregado', data.get('cambio', 0.0)))
            metodo_pago = data.get('metodo_pago', 'efectivo')

            vendedor_id = session.get('usuario_id') or session.get('user_id')
            vendedor_nombre = session.get('usuario') or session.get('username') or 'Cajero'
            num_venta = f"VNT-{int(time.time())}"
            fecha_actual = datetime.now(ZONA_HORARIA_PERU)
            productos_json = json.dumps(items, ensure_ascii=False)

            cursor.execute("""
                INSERT INTO ventas (
                    numero_venta, fecha_venta, cliente_id, cliente_nombre, cliente_documento, 
                    vendedor_id, vendedor_nombre, metodo_pago, subtotal, igv, total, 
                    monto_recibido, cambio_entregado, productos, estado, tenant_id
                ) VALUES (
                    %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, 'completada'::estado_venta_enum, %s
                ) RETURNING id
            """, (
                num_venta, fecha_actual, final_cliente_id, cliente_nombre, cliente_documento,
                vendedor_id, vendedor_nombre, metodo_pago,
                subtotal, igv, total, monto_recibido, cambio, productos_json, tenant_id
            ))

            venta_id = cursor.fetchone()[0]

            for item in items:
                pid = item.get('id')
                cant = int(item.get('cantidad', 1))
                if pid and str(pid).isdigit():
                    cursor.execute("UPDATE productos SET cantidad = GREATEST(COALESCE(cantidad, 0) - %s, 0) WHERE id = %s", (cant, int(pid)))

        conexion.commit()
        conexion.close()

        return jsonify({'success': True, 'venta_id': venta_id, 'ticket_url': f"/venta/ticket/{venta_id}"})

    except Exception as e:
        logger.error(f"❌ ERROR PROCESANDO VENTA: {e}")
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/ventas/scan-push', methods=['POST'])
def scan_push():
    data = request.get_json(silent=True) or {}
    code = data.get('code')
    if code:
        RECENT_SCANS.append({'code': code, 'timestamp': time.time()})
        if len(RECENT_SCANS) > 20:
            RECENT_SCANS.pop(0)
        return jsonify({'success': True, 'code': code})
    return jsonify({'success': False, 'error': 'Código no enviado'}), 400


@app.route('/ventas/scan-poll')
def scan_poll():
    last_time = float(request.args.get('since', 0))
    scans = [s for s in RECENT_SCANS if s['timestamp'] > last_time]
    return jsonify({'success': True, 'scans': scans, 'timestamp': time.time()})


# ==============================================================================
# HISTORIAL DE CLIENTES (DASHBOARD)
# ==============================================================================
@app.route('/historial_clientes', endpoint='historial_clientes')
@app.route('/historial-clientes')
@app.route('/clientes/historial')
def historial_clientes():
    if 'rol' not in session or session['rol'] not in ['admin', 'empleado', 'dueño']:
        flash('Acceso denegado.', 'error')
        return redirect(url_for('dashboard'))
    
    tenant_id = session.get('tenant_id', 1)
    fecha_inicio = request.args.get('fecha_inicio', '').strip()
    fecha_fin = request.args.get('fecha_fin', '').strip()
    busqueda = request.args.get('busqueda', '').strip().lower()
    
    filtros = {'fecha_inicio': fecha_inicio, 'fecha_fin': fecha_fin, 'busqueda': busqueda}
    
    clientes_dict = {}
    id_map = {}
    name_map = {}
    conexion = None

    try:
        conexion = obtener_conexion()
        with conexion.cursor() as cursor:
            col_fecha_ventas = obtener_columna_fecha(cursor)

            # 1. Cargar Clientes registrados en la Base de Datos
            cursor.execute("""
                SELECT 
                    id, 
                    COALESCE(nombre, 'Sin Nombre') AS nombre, 
                    COALESCE(email, 'No registrado') AS email,
                    COALESCE(telefono, '') AS telefono, 
                    COALESCE(direccion, '') AS direccion
                FROM clientes
                WHERE (tenant_id = %s OR tenant_id IS NULL)
                ORDER BY id DESC
            """, (tenant_id,))
            
            for row in cursor.fetchall():
                c_id, c_nom, c_email, c_tel, c_dir = row[0], row[1], row[2], row[3], row[4]
                norm_name = c_nom.strip().lower()
                
                cli_obj = {
                    'id': c_id, 
                    'nombre': c_nom, 
                    'email': c_email, 
                    'telefono': c_tel, 
                    'direccion': c_dir,
                    'fecha_registro': None, 
                    'total_citas': 0, 
                    'total_pedidos': 0,
                    'total_gastado': 0.0, 
                    'ultima_cita': None, 
                    'ultimo_pedido': None
                }
                
                clientes_dict[f"id_{c_id}"] = cli_obj
                id_map[c_id] = cli_obj
                if norm_name:
                    name_map[norm_name] = cli_obj

            # 2. Consultar Ventas y vincular por ID o Nombre
            cursor.execute("""
                SELECT 1 FROM information_schema.columns 
                WHERE table_name = 'ventas' AND column_name = 'cliente_id'
            """)
            has_cliente_id = cursor.fetchone() is not None

            if has_cliente_id:
                query_ventas = f"""
                    SELECT 
                        COALESCE(cliente_id, 0) AS cliente_id,
                        COALESCE(cliente_nombre, 'Cliente General') AS cliente_nombre, 
                        COALESCE(total, 0.0) AS total, 
                        {col_fecha_ventas} AS fecha
                    FROM ventas 
                    WHERE (tenant_id = %s OR tenant_id IS NULL)
                """
            else:
                query_ventas = f"""
                    SELECT 
                        0 AS cliente_id,
                        COALESCE(cliente_nombre, 'Cliente General') AS cliente_nombre, 
                        COALESCE(total, 0.0) AS total, 
                        {col_fecha_ventas} AS fecha
                    FROM ventas 
                    WHERE (tenant_id = %s OR tenant_id IS NULL)
                """

            cursor.execute(query_ventas, (tenant_id,))
            
            for v_row in cursor.fetchall():
                v_cid = int(v_row[0] or 0)
                v_nom = (v_row[1] or 'Cliente General').strip()
                v_tot = float(v_row[2] or 0.0)
                v_fec = str(v_row[3]) if v_row[3] else None
                norm_v_nom = v_nom.lower()

                target_cli = None
                if v_cid and v_cid in id_map:
                    target_cli = id_map[v_cid]
                elif norm_v_nom in name_map:
                    target_cli = name_map[norm_v_nom]
                else:
                    display_name = v_nom if (v_nom and v_nom.strip()) else 'Cliente General'
                    target_cli = {
                        'id': 0, 'nombre': display_name, 'email': 'No registrado', 'telefono': '', 'direccion': '',
                        'fecha_registro': v_fec, 'total_citas': 0, 'total_pedidos': 0,
                        'total_gastado': 0.0, 'ultima_cita': None, 'ultimo_pedido': None
                    }
                    key_virtual = f"name_{norm_v_nom}"
                    clientes_dict[key_virtual] = target_cli
                    name_map[norm_v_nom] = target_cli

                target_cli['total_pedidos'] += 1
                target_cli['total_gastado'] += v_tot
                if v_fec and (not target_cli['ultimo_pedido'] or v_fec > str(target_cli['ultimo_pedido'])):
                    target_cli['ultimo_pedido'] = v_fec

            # 3. Consultar Citas
            cursor.execute("""
                SELECT 
                    COALESCE(cliente_nombre, 'Cliente General') AS cliente_nombre, 
                    COALESCE(precio_total, 0.0) AS total, 
                    fecha 
                FROM citas 
                WHERE (tenant_id = %s OR tenant_id IS NULL)
            """, (tenant_id,))
            
            for c_row in cursor.fetchall():
                c_nom = (c_row[0] or 'Cliente General').strip()
                c_tot = float(c_row[1] or 0.0)
                c_fec = str(c_row[2]) if c_row[2] else None
                norm_c_nom = c_nom.lower()

                target_cli = None
                if norm_c_nom in name_map:
                    target_cli = name_map[norm_c_nom]
                else:
                    display_name = c_nom if (c_nom and c_nom.strip()) else 'Cliente General'
                    target_cli = {
                        'id': 0, 'nombre': display_name, 'email': 'No registrado', 'telefono': '', 'direccion': '',
                        'fecha_registro': c_fec, 'total_citas': 0, 'total_pedidos': 0,
                        'total_gastado': 0.0, 'ultima_cita': None, 'ultimo_pedido': None
                    }
                    key_virtual = f"name_{norm_c_nom}"
                    clientes_dict[key_virtual] = target_cli
                    name_map[norm_c_nom] = target_cli

                target_cli['total_citas'] += 1
                target_cli['total_gastado'] += c_tot
                if c_fec and (not target_cli['ultima_cita'] or c_fec > str(target_cli['ultima_cita'])):
                    target_cli['ultima_cita'] = c_fec

    except Exception as e:
        if conexion:
            conexion.rollback()
        logger.error(f"Error cargando historial clientes: {e}")
    finally:
        if conexion:
            conexion.close()

    historial_lista = []
    total_citas_sum, total_pedidos_sum, total_gastado_sum = 0, 0, 0.0

    for cli in clientes_dict.values():
        if busqueda and (busqueda not in cli['nombre'].lower() and busqueda not in cli['email'].lower()):
            continue

        c_id, c_nom, c_email, c_tel = cli['id'], cli['nombre'], cli['email'], cli['telefono']
        c_freg, c_citas, c_pedidos = cli['fecha_registro'], cli['total_citas'], cli['total_pedidos']
        c_gastado, c_ucita, c_upedido = cli['total_gastado'], cli['ultima_cita'], cli['ultimo_pedido']
        gastado_fmt = round(c_gastado, 2)
        ucita_str = str(c_ucita) if c_ucita else '-'
        upedido_str = str(c_upedido) if c_upedido else '-'
        freg_str = str(c_freg) if c_freg else '-'
        fecha_ult = upedido_str if upedido_str != '-' else (ucita_str if ucita_str != '-' else freg_str)

        obj = {
            'id': c_id, 'cliente_id': c_id, 'id_cliente': c_id,
            'nombre': c_nom, 'cliente': c_nom, 'cliente_nombre': c_nom, 'nombre_cliente': c_nom,
            'email': c_email, 'cliente_email': c_email, 'correo': c_email,
            'telefono': c_tel, 'cliente_telefono': c_tel, 'celular': c_tel,
            'direccion': cli['direccion'], 'cliente_direccion': cli['direccion'],
            'fecha_registro': freg_str, 'fecha': fecha_ult, 'fecha_atencion': fecha_ult,
            'total_citas': c_citas, 'citas': c_citas,
            'total_pedidos': c_pedidos, 'pedidos': c_pedidos,
            'total_atenciones': c_citas + c_pedidos, 'atenciones': c_citas + c_pedidos,
            'total_gastado': gastado_fmt, 'monto_total': gastado_fmt, 'total_monto': gastado_fmt,
            'total_vendido': gastado_fmt, 'monto': gastado_fmt, 'total': gastado_fmt, 'precio': gastado_fmt, 'precio_total': gastado_fmt,
            'ultima_cita': ucita_str, 'ultimo_pedido': upedido_str,
            'servicio': f"{c_citas} citas / {c_pedidos} ventas",
            'servicio_nombre': f"{c_citas} citas / {c_pedidos} ventas",
            'servicio_descripcion': f"Atención integral cliente {c_nom}",
            'comprobante': f"CLI-{c_id}", 'numero_comprobante': f"CLI-{c_id}",
            'estado': 'completado', 'origen': 'Historial Clientes',
            0: c_id, 1: c_nom, 2: c_email, 3: c_tel, 4: cli['direccion'], 5: freg_str,
            6: c_citas, 7: c_pedidos, 8: gastado_fmt, 9: ucita_str, 10: upedido_str
        }
        historial_lista.append(obj)
        total_citas_sum += c_citas
        total_pedidos_sum += c_pedidos
        total_gastado_sum += c_gastado

    historial_lista.sort(key=lambda x: x['total_gastado'], reverse=True)

    total_clientes_cnt = len(historial_lista)
    total_gastado_final = round(total_gastado_sum, 2)
    promedio_ticket_val = round(total_gastado_final / total_clientes_cnt, 2) if total_clientes_cnt > 0 else 0.0

    return render_template(
        'historial_clientes.html', 
        historial=historial_lista, 
        clientes=historial_lista,
        servicios=historial_lista,
        atenciones=historial_lista,
        registros=historial_lista,
        
        total_citas=total_citas_sum, 
        total_pedidos=total_pedidos_sum, 
        total_gastado=total_gastado_final,
        
        total_servicios=total_clientes_cnt,
        total_clientes=total_clientes_cnt,
        total_registros=total_clientes_cnt,
        total_atenciones=total_clientes_cnt,
        
        total_monto=total_gastado_final,
        monto_total=total_gastado_final,
        total_ingresos=total_gastado_final,
        ingresos_recaudados=total_gastado_final,
        
        promedio_ticket=promedio_ticket_val,
        promedio_por_ticket=promedio_ticket_val,
        promedio=promedio_ticket_val,
        
        filtros=filtros
    )


# ==============================================================================
# HISTORIAL DE ASISTENCIA
# ==============================================================================
@app.route('/historial_asistencia', endpoint='historial_asistencia')
@app.route('/historial-asistencia')
@app.route('/asistencia/historial')
def historial_asistencia():
    if 'rol' not in session:
        return redirect(url_for('login'))
        
    tenant_id = session.get('tenant_id', 1)
    usuario_id = obtener_usuario_id_sesion()
    rol = session.get('rol', 'empleado')

    fecha_inicio = request.args.get('fecha_inicio', '').strip()
    fecha_fin = request.args.get('fecha_fin', '').strip()
    busqueda = request.args.get('busqueda', '').strip().lower()

    filtros = {'fecha_inicio': fecha_inicio, 'fecha_fin': fecha_fin, 'busqueda': busqueda}
    registros = []
    
    total_horas_acumuladas = 0.0
    empleados_presentes = 0

    try:
        conexion = obtener_conexion()
        with conexion.cursor() as cursor:
            condiciones = ["(a.tenant_id = %s OR a.tenant_id IS NULL)"]
            params = [tenant_id]

            if rol not in ['admin', 'dueño']:
                condiciones.append("p.usuario_id = %s")
                params.append(usuario_id)

            if fecha_inicio:
                condiciones.append("a.fecha >= %s::date")
                params.append(fecha_inicio)

            if fecha_fin:
                condiciones.append("a.fecha <= %s::date")
                params.append(fecha_fin)

            if busqueda:
                condiciones.append("(LOWER(u.username) LIKE %s OR LOWER(COALESCE(p.cargo, '')) LIKE %s)")
                params.extend([f"%{busqueda}%", f"%{busqueda}%"])

            where_clause = " WHERE " + " AND ".join(condiciones)

            query = f"""
                SELECT 
                    a.id,
                    u.username,
                    COALESCE(p.cargo, 'Empleado') AS cargo,
                    a.fecha,
                    TO_CHAR(a.hora_entrada, 'HH12:MI AM') AS entrada,
                    TO_CHAR(a.hora_salida, 'HH12:MI AM') AS salida,
                    a.hora_entrada AS raw_entrada,
                    a.hora_salida AS raw_salida
                FROM asistencia a
                INNER JOIN personal p ON a.personal_id = p.id
                INNER JOIN usuarios u ON p.usuario_id = u.id
                {where_clause}
                ORDER BY a.fecha DESC, a.hora_entrada DESC
                LIMIT 300
            """
            cursor.execute(query, tuple(params))
            rows = cursor.fetchall()

            for r in rows:
                a_id, u_nom, p_cargo, a_fec = r[0], r[1] or 'Empleado', r[2] or 'Empleado', str(r[3])
                h_ent, h_sal = r[4] or '--:--', r[5] or 'En turno'
                raw_ent, raw_sal = r[6], r[7]

                horas_trab = 0.0
                if raw_ent and raw_sal:
                    try:
                        dummy_d = datetime.now(ZONA_HORARIA_PERU).date()
                        dt_ent = datetime.combine(dummy_d, raw_ent)
                        dt_sal = datetime.combine(dummy_d, raw_sal)
                        delta = dt_sal - dt_ent
                        horas_trab = max(0.0, round(delta.total_seconds() / 3600, 1))
                    except Exception:
                        horas_trab = 0.0

                total_horas_acumuladas += horas_trab
                if not raw_sal:
                    empleados_presentes += 1

                estado_str = 'completado' if raw_sal else 'presente'

                item = {
                    'id': a_id, 'usuario': u_nom, 'nombre': u_nom, 'username': u_nom, 'cargo': p_cargo,
                    'fecha': a_fec, 'hora_entrada': h_ent, 'entrada': h_ent, 'hora_salida': h_sal, 'salida': h_sal,
                    'horas_trabajadas': horas_trab, 'estado': estado_str,
                    0: a_id, 1: u_nom, 2: p_cargo, 3: a_fec, 4: h_ent, 5: h_sal, 6: estado_str, 7: horas_trab
                }
                registros.append(item)

        conexion.close()
    except Exception as e:
        logger.error(f"Error cargando historial de asistencia: {e}")

    return render_template(
        'historial_asistencia.html', 
        historial=registros, 
        registros=registros,
        total_asistencias=len(registros),
        total_horas=round(total_horas_acumuladas, 1),
        empleados_presentes=empleados_presentes,
        filtros=filtros
    )


# ==============================================================================
# HISTORIAL DE COMPRAS
# ==============================================================================
@app.route('/historial_compras', endpoint='historial_compras')
@app.route('/historial-compras')
@app.route('/compras/historial')
def historial_compras():
    if 'rol' not in session or session['rol'] not in ['admin', 'dueño', 'empleado']:
        flash('Acceso denegado.', 'error')
        return redirect(url_for('dashboard'))
    
    tenant_id = session.get('tenant_id', 1)
    fecha_inicio = request.args.get('fecha_inicio', '').strip()
    fecha_fin = request.args.get('fecha_fin', '').strip()
    busqueda = request.args.get('busqueda', '').strip().lower()

    filtros = {'fecha_inicio': fecha_inicio, 'fecha_fin': fecha_fin, 'busqueda': busqueda}
    compras_items = []
    
    unidades_totales = 0
    monto_total_acumulado = 0.0

    try:
        conexion = obtener_conexion()
        with conexion.cursor() as cursor:
            col_fecha = obtener_columna_fecha(cursor)
            
            condiciones = ["(v.tenant_id = %s OR v.tenant_id IS NULL)"]
            params = [tenant_id]

            if fecha_inicio:
                condiciones.append(f"DATE(v.{col_fecha}::text) >= %s::date")
                params.append(fecha_inicio)

            if fecha_fin:
                condiciones.append(f"DATE(v.{col_fecha}::text) <= %s::date")
                params.append(fecha_fin)

            where_clause = " WHERE " + " AND ".join(condiciones)

            query = f"""
                SELECT 
                    v.id,
                    v.numero_venta,
                    TO_CHAR(v.{col_fecha}, 'DD/MM/YYYY HH12:MI AM') AS fecha_fmt,
                    COALESCE(v.cliente_nombre, 'Cliente General') AS cliente,
                    COALESCE(v.metodo_pago, 'efectivo') AS pago,
                    v.productos
                FROM ventas v
                {where_clause}
                ORDER BY v.{col_fecha} DESC
                LIMIT 300
            """
            cursor.execute(query, tuple(params))
            rows = cursor.fetchall()

            for r in rows:
                v_id, v_num, v_fec, v_cli, v_pago, raw_prod = r[0], r[1] or 'N/A', r[2] or '', r[3], r[4], r[5]
                
                if not raw_prod:
                    continue

                try:
                    items = json.loads(raw_prod) if isinstance(raw_prod, str) else raw_prod
                    if not isinstance(items, list):
                        continue

                    for item in items:
                        p_id = item.get('id')
                        p_nombre = item.get('nombre') or item.get('title') or 'Producto'
                        
                        if str(p_id).startswith('service_') or item.get('type') == 'service':
                            continue

                        cant = int(item.get('cantidad', item.get('qty', 1)))
                        precio = float(item.get('precio', item.get('precio_unitario', 0.0)))
                        subtotal = cant * precio
                        subtotal_fmt = round(subtotal, 2)

                        if busqueda and (busqueda not in p_nombre.lower() and busqueda not in v_cli.lower() and busqueda not in str(v_num).lower()):
                            continue

                        unidades_totales += cant
                        monto_total_acumulado += subtotal

                        compras_items.append({
                            'venta_id': v_id,
                            'numero_comprobante': v_num,
                            'comprobante': v_num,
                            'numero_venta': v_num,
                            'fecha': v_fec,
                            'cliente_nombre': v_cli,
                            'cliente': v_cli,
                            'producto_id': p_id,
                            'producto_nombre': p_nombre,
                            'producto': p_nombre,
                            'nombre': p_nombre,
                            'cantidad': cant,
                            'unidades': cant,
                            'precio_unitario': precio,
                            'precio': precio,
                            'subtotal': subtotal_fmt,
                            'total': subtotal_fmt,
                            'monto': subtotal_fmt,
                            'monto_total': subtotal_fmt,
                            'total_vendido': subtotal_fmt,
                            'metodo_pago': v_pago,
                            'pago': v_pago,
                            0: v_id, 1: v_num, 2: v_fec, 3: v_cli, 4: p_nombre, 5: cant, 6: precio, 7: subtotal_fmt, 8: v_pago
                        })
                except Exception as ex_json:
                    logger.warning(f"Error procesando JSON de productos para venta {v_id}: {ex_json}")

        conexion.close()
    except Exception as e:
        logger.error(f"Error consultando historial de compras: {e}")

    total_acumulado_fmt = round(monto_total_acumulado, 2)

    estadisticas = {
        'total_compras': len(compras_items),
        'total_items': len(compras_items),
        'unidades_totales': unidades_totales,
        'total_unidades': unidades_totales,
        'monto_total': total_acumulado_fmt,
        'total_monto': total_acumulado_fmt,
        'total_vendido': total_acumulado_fmt
    }

    return render_template('historial_compras.html', compras=compras_items, estadisticas=estadisticas, filtros=filtros)


# ==============================================================================
# HISTORIAL DE SERVICIOS
# ==============================================================================
@app.route('/historial_servicios', endpoint='historial_servicios')
@app.route('/historial-servicios')
@app.route('/servicios/historial')
def historial_servicios():
    if 'rol' not in session or session['rol'] not in ['admin', 'empleado', 'dueño']:
        flash('Acceso denegado.', 'error')
        return redirect(url_for('dashboard'))
    
    tenant_id = session.get('tenant_id', 1)
    fecha_inicio = request.args.get('fecha_inicio', '').strip()
    fecha_fin = request.args.get('fecha_fin', '').strip()
    busqueda = request.args.get('busqueda', '').strip().lower()

    filtros = {'fecha_inicio': fecha_inicio, 'fecha_fin': fecha_fin, 'busqueda': busqueda}
    servicios_historial = []
    
    monto_total_servicios = 0.0
    servicios_completados = 0

    try:
        conexion = obtener_conexion()
        with conexion.cursor() as cursor:
            cond_citas = ["(c.tenant_id = %s OR c.tenant_id IS NULL)"]
            params_citas = [tenant_id]

            if fecha_inicio:
                cond_citas.append("c.fecha >= %s::date")
                params_citas.append(fecha_inicio)

            if fecha_fin:
                cond_citas.append("c.fecha <= %s::date")
                params_citas.append(fecha_fin)

            where_citas = " WHERE " + " AND ".join(cond_citas)

            query_citas = f"""
                SELECT 
                    c.id,
                    TO_CHAR(c.fecha, 'DD/MM/YYYY') AS fecha_fmt,
                    TO_CHAR(c.hora, 'HH12:MI AM') AS hora_fmt,
                    COALESCE(c.cliente_nombre, 'Cliente General') AS cliente,
                    COALESCE(c.mascota_nombre, '-') AS mascota,
                    COALESCE(s.nombre, 'Servicio General') AS servicio,
                    COALESCE(c.precio_total, 0.0) AS precio,
                    COALESCE(c.estado::text, 'pendiente') AS estado
                FROM citas c
                LEFT JOIN servicios s ON c.servicio_id = s.id
                {where_citas}
                ORDER BY c.fecha DESC, c.hora DESC
            """
            cursor.execute(query_citas, tuple(params_citas))
            
            for r in cursor.fetchall():
                c_id, c_fec, c_hor, c_cli, c_masc, s_nom, s_prec, s_est = r[0], r[1], r[2], r[3], r[4], r[5], float(r[6] or 0), str(r[7]).lower()

                if busqueda and (busqueda not in s_nom.lower() and busqueda not in c_cli.lower() and busqueda not in c_masc.lower()):
                    continue

                monto_total_servicios += s_prec
                if s_est in ['completada', 'completado', 'confirmada']:
                    servicios_completados += 1

                prec_fmt = round(s_prec, 2)

                servicios_historial.append({
                    'id': f"CITA-{c_id}",
                    'fecha': c_fec,
                    'hora': c_hor,
                    'cliente_nombre': c_cli,
                    'cliente': c_cli,
                    'mascota_nombre': c_masc,
                    'mascota': c_masc,
                    'servicio_nombre': s_nom,
                    'servicio': s_nom,
                    'nombre': s_nom,
                    'precio': prec_fmt,
                    'precio_total': prec_fmt,
                    'subtotal': prec_fmt,
                    'monto': prec_fmt,
                    'monto_total': prec_fmt,
                    'total': prec_fmt,
                    'total_vendido': prec_fmt,
                    'estado': s_est,
                    'origen': 'Cita Agendada',
                    0: f"CITA-{c_id}", 1: c_fec, 2: c_cli, 3: c_masc, 4: s_nom, 5: prec_fmt, 6: s_est, 7: 'Cita Agendada'
                })

            col_fecha = obtener_columna_fecha(cursor)
            cond_v = ["(v.tenant_id = %s OR v.tenant_id IS NULL)"]
            params_v = [tenant_id]

            if fecha_inicio:
                cond_v.append(f"DATE(v.{col_fecha}::text) >= %s::date")
                params_v.append(fecha_inicio)

            if fecha_fin:
                cond_v.append(f"DATE(v.{col_fecha}::text) <= %s::date")
                params_v.append(fecha_fin)

            where_v = " WHERE " + " AND ".join(cond_v)

            query_v = f"""
                SELECT 
                    v.id,
                    TO_CHAR(v.{col_fecha}, 'DD/MM/YYYY') AS fecha_fmt,
                    COALESCE(v.cliente_nombre, 'Cliente General') AS cliente,
                    v.productos
                FROM ventas v
                {where_v}
            """
            cursor.execute(query_v, tuple(params_v))
            
            for r in cursor.fetchall():
                v_id, v_fec, v_cli, raw_prod = r[0], r[1], r[2], r[3]
                if not raw_prod:
                    continue

                try:
                    items = json.loads(raw_prod) if isinstance(raw_prod, str) else raw_prod
                    if not isinstance(items, list):
                        continue

                    for item in items:
                        p_id = item.get('id')
                        p_nom = item.get('nombre') or item.get('title') or 'Servicio POS'

                        if str(p_id).startswith('service_') or item.get('type') == 'service':
                            prec = float(item.get('precio', item.get('precio_unitario', 0.0)))
                            cant = int(item.get('cantidad', item.get('qty', 1)))
                            total_s = prec * cant
                            total_s_fmt = round(total_s, 2)

                            if busqueda and (busqueda not in p_nom.lower() and busqueda not in v_cli.lower()):
                                continue

                            monto_total_servicios += total_s
                            servicios_completados += 1

                            servicios_historial.append({
                                'id': f"POS-{v_id}",
                                'fecha': v_fec,
                                'hora': '--:--',
                                'cliente_nombre': v_cli,
                                'cliente': v_cli,
                                'mascota_nombre': '-',
                                'mascota': '-',
                                'servicio_nombre': p_nom,
                                'servicio': p_nom,
                                'nombre': p_nom,
                                'precio': total_s_fmt,
                                'precio_total': total_s_fmt,
                                'subtotal': total_s_fmt,
                                'monto': total_s_fmt,
                                'monto_total': total_s_fmt,
                                'total': total_s_fmt,
                                'total_vendido': total_s_fmt,
                                'estado': 'completado',
                                'origen': 'Punto de Venta',
                                0: f"POS-{v_id}", 1: v_fec, 2: v_cli, 3: '-', 4: p_nom, 5: total_s_fmt, 6: 'completado', 7: 'Punto de Venta'
                            })
                except Exception:
                    pass

        conexion.close()
    except Exception as e:
        logger.error(f"Error consultando historial de servicios: {e}")

    total_serv_fmt = round(monto_total_servicios, 2)

    estadisticas = {
        'total_servicios': len(servicios_historial),
        'total_items': len(servicios_historial),
        'monto_total': total_serv_fmt,
        'total_monto': total_serv_fmt,
        'total_vendido': total_serv_fmt,
        'servicios_completados': servicios_completados
    }

    return render_template('historial_servicios.html', servicios=servicios_historial, estadisticas=estadisticas, filtros=filtros)


# ==============================================================================
# HISTORIAL DE VENTAS GENERALES
# ==============================================================================
@app.route('/historial_ventas', endpoint='historial_ventas')
@app.route('/historial-ventas')
@app.route('/ventas/historial')
def historial_ventas():
    if 'rol' not in session or session['rol'] not in ['admin', 'empleado', 'dueño']:
        flash('Acceso denegado.', 'error')
        return redirect(url_for('dashboard'))
    
    tenant_id = session.get('tenant_id', 1)
    fecha_inicio = request.args.get('fecha_inicio', '').strip()
    fecha_fin = request.args.get('fecha_fin', '').strip()
    busqueda = request.args.get('busqueda', '').strip().lower()
    
    filtros = {'fecha_inicio': fecha_inicio, 'fecha_fin': fecha_fin, 'busqueda': busqueda}
    ventas_lista = []
    conexion = None

    try:
        conexion = obtener_conexion()
        with conexion.cursor() as cursor:
            col_fecha = obtener_columna_fecha(cursor)

            condiciones = ["(v.tenant_id = %s OR v.tenant_id IS NULL)"]
            params = [tenant_id]

            if fecha_inicio:
                condiciones.append(f"DATE(v.{col_fecha}::text) >= %s::date")
                params.append(fecha_inicio)

            if fecha_fin:
                condiciones.append(f"DATE(v.{col_fecha}::text) <= %s::date")
                params.append(fecha_fin)

            if busqueda:
                condiciones.append("(COALESCE(v.cliente_nombre, '') ILIKE %s OR COALESCE(v.numero_venta, '') ILIKE %s)")
                params.extend([f"%{busqueda}%", f"%{busqueda}%"])

            where_clause = " WHERE " + " AND ".join(condiciones)

            query = f"""
                SELECT 
                    v.id,
                    COALESCE(v.numero_venta, 'N/A') AS comprobante,
                    TO_CHAR(v.{col_fecha}, 'DD/MM/YYYY HH12:MI AM') AS fecha,
                    COALESCE(v.cliente_nombre, 'Cliente General') AS cliente,
                    'Venta General' AS servicio,
                    COALESCE(v.vendedor_nombre, 'Atendido') AS personal,
                    COALESCE(v.metodo_pago, 'efectivo') AS pago,
                    COALESCE(v.total, 0.00) AS total,
                    COALESCE(v.estado::text, 'completado') AS estado
                FROM ventas v
                {where_clause}
                ORDER BY v.{col_fecha} DESC
                LIMIT 200
            """

            cursor.execute(query, tuple(params))
            rows = cursor.fetchall()

            for r in rows:
                v_id, v_num, v_fec, v_cli, v_serv, v_pers, v_pago, v_tot, v_est = r[0], r[1], r[2], r[3], r[4], r[5], r[6], float(r[7] or 0), str(r[8]).lower()

                ventas_lista.append({
                    'id': v_id,
                    'numero_comprobante': v_num,
                    'comprobante': v_num,
                    'fecha': str(v_fec),
                    'cliente_nombre': v_cli,
                    'servicio_descripcion': v_serv,
                    'personal': v_pers,
                    'metodo_pago': v_pago,
                    'total': v_tot,
                    'total_vendido': v_tot,
                    'estado': v_est,
                    0: v_id, 1: v_num, 2: str(v_fec), 3: v_cli, 4: v_serv, 5: v_pers, 6: v_pago, 7: v_tot, 8: v_est
                })

    except Exception as e:
        if conexion:
            conexion.rollback()
        logger.error(f"Error en historial_ventas: {e}")
    finally:
        if conexion:
            conexion.close()

    try:
        return render_template('historial_ventas.html', ventas=ventas_lista, servicios=ventas_lista, filtros=filtros)
    except Exception:
        return render_template('historial_clientes.html', ventas=ventas_lista, servicios=ventas_lista, filtros=filtros)


# ─── VISTA Y GENERACIÓN DE TICKET DE VENTA ───────────────────────────────────

@app.route('/venta/ticket/<int:venta_id>', endpoint='ticket_venta')
@app.route('/ventas/ticket/<int:venta_id>')
def ticket_venta(venta_id):
    if 'rol' not in session:
        return redirect(url_for('login'))

    tenant_id = session.get('tenant_id', 1)
    venta = None
    items = []

    try:
        conexion = obtener_conexion()
        with conexion.cursor() as cursor:
            col_fecha = obtener_columna_fecha(cursor)
            cursor.execute(f"""
                SELECT 
                    v.id,
                    v.numero_venta,
                    TO_CHAR(v.{col_fecha}, 'DD/MM/YYYY HH12:MI AM') AS fecha,
                    COALESCE(v.cliente_nombre, 'Cliente General') AS cliente,
                    COALESCE(v.cliente_documento, '-') AS doc,
                    COALESCE(v.vendedor_nombre, 'Cajero') AS vendedor,
                    COALESCE(v.metodo_pago, 'Efectivo') AS pago,
                    COALESCE(v.subtotal, 0.00) AS subtotal,
                    COALESCE(v.igv, 0.00) AS igv,
                    COALESCE(v.total, 0.00) AS total,
                    COALESCE(v.monto_recibido, 0.00) AS recibido,
                    COALESCE(v.cambio_entregado, 0.00) AS cambio,
                    COALESCE(v.estado::text, 'completada') AS estado,
                    v.productos
                FROM ventas v
                WHERE v.id = %s AND (v.tenant_id = %s OR v.tenant_id IS NULL)
            """, (venta_id, tenant_id))
            
            r = cursor.fetchone()
            if r:
                subtotal_f = float(r[7]) if r[7] is not None else 0.0
                igv_f = float(r[8]) if r[8] is not None else 0.0
                total_f = float(r[9]) if r[9] is not None else 0.0
                recibido_f = float(r[10]) if r[10] is not None else 0.0
                cambio_f = float(r[11]) if r[11] is not None else 0.0

                venta = {
                    'id': r[0], 'numero_venta': r[1], 'fecha_venta': r[2], 'cliente_nombre': r[3],
                    'cliente_documento': r[4], 'vendedor_nombre': r[5], 'metodo_pago': r[6],
                    'subtotal': subtotal_f, 'igv': igv_f, 'total': total_f,
                    'monto_recibido': recibido_f, 'cambio_entregado': cambio_f, 'estado': str(r[12]).lower(),
                    0: r[0], 1: r[1], 2: r[2], 3: r[3], 4: r[4], 5: r[5], 6: r[6],
                    7: subtotal_f, 8: igv_f, 9: total_f, 10: recibido_f, 11: cambio_f, 12: str(r[12]).lower()
                }

                raw_productos = r[13]
                if raw_productos:
                    try:
                        p_list = json.loads(raw_productos) if isinstance(raw_productos, str) else raw_productos
                        for p in p_list:
                            cant = int(p.get('cantidad', p.get('qty', 1)))
                            pu = float(p.get('precio', p.get('precio_unitario', 0.0)))
                            items.append({
                                'nombre': p.get('nombre', p.get('title', 'Producto / Servicio')),
                                'cantidad': cant, 'precio_unitario': pu, 'subtotal': cant * pu
                            })
                    except Exception as ex_json:
                        logger.warning(f"No se pudo decodificar JSON de productos para venta {venta_id}: {ex_json}")

        conexion.close()
    except Exception as e:
        logger.error(f"Error generando ticket para venta ID={venta_id}: {e}")

    if not venta:
        flash('La venta solicitada no existe.', 'error')
        return redirect(url_for('historial_clientes'))

    return render_template('ticket_venta.html', venta=venta, items=items, momento_actual=datetime.now(ZONA_HORARIA_PERU))


# ─── INVENTARIO & ALMACÉN ───────────────────────────────────────────────────

@app.route('/productos')
def productos():
    if 'rol' not in session or session['rol'] not in ['admin', 'empleado', 'dueño']:
        flash('Acceso denegado.', 'error')
        return redirect(url_for('dashboard'))
    
    tenant_id = session.get('tenant_id', 1)
    if hasattr(productos_controlador, 'obtener_productos'):
        try:
            lista = productos_controlador.obtener_productos(tenant_id)
        except TypeError:
            lista = productos_controlador.obtener_productos()
    else:
        lista = []
    return render_template('productos.html', productos=lista)


@app.route('/almacen', methods=['GET', 'POST'])
@app.route('/productos/agregar', methods=['POST'], endpoint='agregar_producto')
@app.route('/productos/editar', methods=['POST'], endpoint='editar_producto')
def almacen():
    if 'rol' not in session or session['rol'] not in ['admin', 'empleado', 'dueño']:
        flash('Acceso denegado.', 'error')
        return redirect(url_for('dashboard'))

    tenant_id = session.get('tenant_id', 1)

    if request.method == 'POST':
        try:
            raw_id = (request.form.get('id') or request.form.get('producto_id') or 
                      request.form.get('edit_id') or request.form.get('id_producto') or '').strip()
            
            producto_id = int(raw_id) if raw_id and raw_id.isdigit() else None
            nombre = request.form.get('nombre', '').strip()
            codigo_barra = request.form.get('codigo_barra', '').strip() or None
            
            tipo = (request.form.get('tipo') or request.form.get('tipo_producto') or 
                    request.form.get('tipo_item') or request.form.get('edit_tipo') or '').strip().lower()

            precio_raw = request.form.get('precio', '').strip()
            precio = float(precio_raw) if precio_raw else 0.0

            cantidad_raw = request.form.get('stock', request.form.get('cantidad', '')).strip()
            cantidad = int(cantidad_raw) if cantidad_raw else 0

            stk_min_raw = (request.form.get('stock_minimo') or request.form.get('cant_min') or request.form.get('stock_min') or '').strip()
            stock_minimo = int(stk_min_raw) if stk_min_raw and stk_min_raw.isdigit() else 5

            conexion = obtener_conexion()
            with conexion.cursor() as cursor:
                cursor.execute("ALTER TABLE productos ADD COLUMN IF NOT EXISTS stock_minimo INT DEFAULT 5")
                cursor.execute("ALTER TABLE productos ADD COLUMN IF NOT EXISTS fecha_vencimiento DATE")
                cursor.execute("ALTER TABLE productos ADD COLUMN IF NOT EXISTS imagen VARCHAR(255)")
                cursor.execute("ALTER TABLE productos ADD COLUMN IF NOT EXISTS tenant_id INT DEFAULT 1")

                if codigo_barra:
                    if producto_id:
                        cursor.execute("SELECT nombre FROM productos WHERE codigo_barra = %s AND id != %s AND (tenant_id = %s OR tenant_id IS NULL)", (codigo_barra, producto_id, tenant_id))
                    else:
                        cursor.execute("SELECT nombre FROM productos WHERE codigo_barra = %s AND (tenant_id = %s OR tenant_id IS NULL)", (codigo_barra, tenant_id))
                    
                    existente = cursor.fetchone()
                    if existente:
                        flash(f'No se puede guardar: El código de barras "{codigo_barra}" pertenece a "{existente[0]}".', 'error')
                        conexion.close()
                        return redirect(url_for('almacen'))

                if producto_id:
                    if not tipo:
                        cursor.execute("SELECT tipo::text FROM productos WHERE id = %s AND (tenant_id = %s OR tenant_id IS NULL)", (producto_id, tenant_id))
                        res_tipo = cursor.fetchone()
                        if res_tipo and res_tipo[0]:
                            tipo = str(res_tipo[0]).strip().lower()
                    if not tipo:
                        tipo = 'stock'

                    cursor.execute("""
                        UPDATE productos 
                        SET nombre = %s, codigo_barra = %s, tipo = %s, cantidad = %s, precio = %s, stock_minimo = %s
                        WHERE id = %s AND (tenant_id = %s OR tenant_id IS NULL)
                    """, (nombre, codigo_barra, tipo, cantidad, precio, stock_minimo, producto_id, tenant_id))
                    
                    mensaje = f'Producto "{nombre}" actualizado correctamente.' if cursor.rowcount > 0 else f'Producto ID {producto_id} no encontrado.'
                    categoria = 'success' if cursor.rowcount > 0 else 'warning'
                else:
                    if not tipo:
                        tipo = 'stock'

                    cursor.execute("""
                        INSERT INTO productos (nombre, codigo_barra, tipo, cantidad, precio, stock_minimo, tenant_id)
                        VALUES (%s, %s, %s, %s, %s, %s, %s)
                    """, (nombre, codigo_barra, tipo, cantidad, precio, stock_minimo, tenant_id))
                    mensaje = f'Producto "{nombre}" registrado correctamente.'
                    categoria = 'success'

            conexion.commit()
            conexion.close()
            flash(mensaje, categoria)
        except Exception as e:
            logger.error(f"Error procesando producto en almacén: {e}")
            flash(f'Error al procesar producto: {e}', 'error')
        return redirect(url_for('almacen'))

    tipo_filtro = request.args.get('tipo', '').strip().lower()
    productos_lista = []

    try:
        conexion = obtener_conexion()
        with conexion.cursor() as cursor:
            cursor.execute("ALTER TABLE productos ADD COLUMN IF NOT EXISTS stock_minimo INT DEFAULT 5")
            cursor.execute("ALTER TABLE productos ADD COLUMN IF NOT EXISTS fecha_vencimiento DATE")
            cursor.execute("ALTER TABLE productos ADD COLUMN IF NOT EXISTS imagen VARCHAR(255)")
            cursor.execute("ALTER TABLE productos ADD COLUMN IF NOT EXISTS tenant_id INT DEFAULT 1")

            sql_query = "SELECT id, nombre, COALESCE(tipo::text, 'stock'), cantidad, precio, COALESCE(stock_minimo, 5), fecha_vencimiento, imagen, codigo_barra FROM productos WHERE (tenant_id = %s OR tenant_id IS NULL)"

            if tipo_filtro == 'stock':
                sql_query += " AND LOWER(COALESCE(tipo::text, 'stock')) = 'stock' "
            elif tipo_filtro == 'venta':
                sql_query += " AND LOWER(tipo::text) = 'venta' "

            sql_query += " ORDER BY id ASC "

            cursor.execute(sql_query, (tenant_id,))
            for r in cursor.fetchall():
                p_id, nombre = r[0], r[1] or ''
                tipo = str(r[2]).strip().lower() if r[2] is not None else 'stock'
                cantidad = int(r[3]) if r[3] is not None else 0
                precio = float(r[4]) if r[4] is not None else 0.0
                stock_minimo = int(r[5]) if r[5] is not None else 5
                fecha_venc, imagen, codigo_barra = r[6], r[7] or '', r[8] or ''

                productos_lista.append({
                    'id': p_id, 'nombre': nombre, 'tipo': tipo, 'cantidad': cantidad, 'stock': cantidad,
                    'precio': precio, 'stock_minimo': stock_minimo, 'fecha_vencimiento': fecha_venc,
                    'imagen': imagen, 'codigo_barra': codigo_barra,
                    0: p_id, 1: nombre, 2: tipo, 3: cantidad, 4: precio, 5: stock_minimo, 6: fecha_venc, 7: imagen, 8: codigo_barra
                })
        conexion.close()
    except Exception as e:
        logger.error(f"Error consultando productos en almacén: {e}")

    return render_template('almacen.html', productos=productos_lista, tipo_filtro=tipo_filtro)


@app.route('/producto/<int:producto_id>')
def ver_producto(producto_id):
    if 'rol' not in session or session['rol'] not in ['admin', 'empleado', 'dueño']:
        flash('Acceso denegado.', 'error')
        return redirect(url_for('dashboard'))

    tenant_id = session.get('tenant_id', 1)

    if hasattr(productos_controlador, 'obtener_producto_por_id'):
        try:
            producto = productos_controlador.obtener_producto_por_id(producto_id, tenant_id)
        except TypeError:
            producto = productos_controlador.obtener_producto_por_id(producto_id)
    else:
        producto = None

    if not producto:
        flash('Producto no encontrado.', 'error')
        return redirect(url_for('productos'))
    return render_template('ver_producto.html', producto=producto)


# ─── MÓDULO DE SERVICIOS COMPLETO ─────────────────────────────────────────────

@app.template_filter('moneda')
def formato_moneda_filter(val):
    try:
        return f"S/ {float(val):.2f}"
    except (ValueError, TypeError):
        return "S/ 0.00"


@app.route('/servicios', methods=['GET'], endpoint='servicios')
@app.route('/gestion_servicios', methods=['GET'], endpoint='gestion_servicios')
def servicios():
    if 'rol' not in session or session['rol'] not in ['admin', 'empleado', 'dueño']:
        flash('Acceso denegado.', 'error')
        return redirect(url_for('dashboard'))
    
    tenant_id = session.get('tenant_id', 1)
    servicios_lista = []

    try:
        conexion = obtener_conexion()
        with conexion.cursor() as cursor:
            cursor.execute("""
                SELECT id, nombre, COALESCE(descripcion, '') AS descripcion, 
                       COALESCE(precio, 0.00) AS precio, COALESCE(duracion, 30) AS duracion, 
                       COALESCE(activo, true) AS activo
                FROM servicios
                WHERE (tenant_id = %s OR tenant_id IS NULL)
                ORDER BY nombre ASC
            """, (tenant_id,))
            
            for s in cursor.fetchall():
                s_id, s_nom, s_desc = s[0], s[1], s[2]
                s_prec = float(s[3]) if s[3] is not None else 0.0
                s_dur = int(s[4]) if s[4] is not None else 30
                s_act = bool(s[5])

                servicios_lista.append({
                    'id': s_id, 'nombre': str(s_nom or ''), 'descripcion': str(s_desc or ''),
                    'precio': s_prec, 'duracion': s_dur, 'duracion_minutos': s_dur, 'activo': s_act,
                    0: s_id, 1: str(s_nom or ''), 2: str(s_desc or ''), 3: s_prec, 4: s_dur, 5: s_act
                })
        conexion.close()
    except Exception as e:
        logger.error(f"Error consultando servicios: {e}")

    return render_template('servicios.html', servicios=servicios_lista)


@app.route('/servicios/agregar', methods=['POST'])
def agregar_servicio():
    if session.get('rol') not in ['admin', 'dueño']:
        return jsonify({'success': False, 'error': 'No autorizado'}), 403
    try:
        tenant_id = session.get('tenant_id', 1)
        nombre = request.form.get('nombre', '').strip()
        descripcion = request.form.get('descripcion', '').strip()
        precio = float(request.form.get('precio', 0.0) or 0.0)
        duracion = int(request.form.get('duracion', 30) or 30)

        if not nombre:
            flash('El nombre del servicio es obligatorio.', 'warning')
            return redirect(url_for('servicios'))

        conexion = obtener_conexion()
        with conexion.cursor() as cursor:
            cursor.execute("""
                INSERT INTO servicios (nombre, descripcion, precio, duracion, activo, tenant_id)
                VALUES (%s, %s, %s, %s, true, %s)
            """, (nombre, descripcion, precio, duracion, tenant_id))
        conexion.commit()
        conexion.close()

        flash('Servicio registrado exitosamente.', 'success')
        return redirect(url_for('servicios'))
    except Exception as e:
        logger.error(f"Error al agregar servicio: {e}")
        flash(f'Error al registrar servicio: {e}', 'error')
        return redirect(url_for('servicios'))


@app.route('/servicios/editar', methods=['POST'])
def editar_servicio():
    if session.get('rol') not in ['admin', 'dueño']:
        return jsonify({'success': False, 'error': 'No autorizado'}), 403
    try:
        tenant_id = session.get('tenant_id', 1)
        raw_id = request.form.get('servicio_id') or request.form.get('id')
        servicio_id = int(raw_id) if raw_id and str(raw_id).isdigit() else None
        
        nombre = request.form.get('nombre', '').strip()
        descripcion = request.form.get('descripcion', '').strip()
        precio = float(request.form.get('precio', 0.0) or 0.0)
        duracion = int(request.form.get('duracion', 30) or 30)

        if not servicio_id or not nombre:
            flash('Datos incompletos para actualizar el servicio.', 'error')
            return redirect(url_for('servicios'))

        conexion = obtener_conexion()
        with conexion.cursor() as cursor:
            cursor.execute("""
                UPDATE servicios 
                SET nombre = %s, descripcion = %s, precio = %s, duracion = %s
                WHERE id = %s AND (tenant_id = %s OR tenant_id IS NULL)
            """, (nombre, descripcion, precio, duracion, servicio_id, tenant_id))
        conexion.commit()
        conexion.close()

        flash('Servicio actualizado exitosamente.', 'success')
        return redirect(url_for('servicios'))
    except Exception as e:
        logger.error(f"Error al editar servicio: {e}")
        flash(f'Error al actualizar servicio: {e}', 'error')
        return redirect(url_for('servicios'))


@app.route('/servicios/cambiar-estado', methods=['POST'])
def cambiar_estado_servicio():
    if session.get('rol') not in ['admin', 'dueño']:
        return jsonify({'success': False, 'error': 'No autorizado'}), 403
    try:
        tenant_id = session.get('tenant_id', 1)
        data = request.get_json(silent=True) or {}
        servicio_id = data.get('servicio_id') or data.get('id')
        activo = data.get('activo', True)
        
        if not servicio_id:
            return jsonify({'success': False, 'error': 'ID no especificado'}), 400

        conexion = obtener_conexion()
        with conexion.cursor() as cursor:
            cursor.execute("""
                UPDATE servicios 
                SET activo = %s 
                WHERE id = %s AND (tenant_id = %s OR tenant_id IS NULL)
            """, (bool(activo), int(servicio_id), tenant_id))
        conexion.commit()
        conexion.close()
        
        return jsonify({'success': True, 'message': 'Estado actualizado.'})
    except Exception as e:
        logger.error(f"Error al cambiar estado del servicio: {e}")
        return jsonify({'success': False, 'error': str(e)}), 500


# ─── CLIENTES & MASCOTAS ─────────────────────────────────────────────────────

@app.route('/clientes', endpoint='clientes')
@app.route('/clientes/lista', endpoint='listar_clientes')
def clientes():
    if 'rol' not in session or session['rol'] not in ['admin', 'empleado', 'dueño']:
        flash('Acceso denegado.', 'error')
        return redirect(url_for('dashboard'))
    
    tenant_id = session.get('tenant_id', 1)
    clientes_lista = []
    
    try:
        conexion = obtener_conexion()
        with conexion.cursor() as cursor:
            cursor.execute("""
                SELECT id, COALESCE(nombre, '') AS nombre, COALESCE(email, '') AS email, 
                       COALESCE(telefono, '') AS telefono, COALESCE(direccion, '') AS direccion, 
                       COALESCE(documento, '') AS documento, COALESCE(activo, true) AS activo
                FROM clientes
                WHERE (tenant_id = %s OR tenant_id IS NULL)
                ORDER BY id DESC
            """, (tenant_id,))
            
            for r in cursor.fetchall():
                c_id, c_nom, c_email, c_tel, c_dir, c_doc, c_act = r[0], r[1], r[2], r[3], r[4], r[5], r[6]
                clientes_lista.append({
                    'id': c_id, 'nombre': str(c_nom or ''), 'email': str(c_email or ''),
                    'telefono': str(c_tel or ''), 'direccion': str(c_dir or ''),
                    'documento': str(c_doc or ''), 'activo': bool(c_act),
                    0: c_id, 1: str(c_nom or ''), 2: str(c_email or ''), 
                    3: str(c_tel or ''), 4: str(c_dir or ''), 5: str(c_doc or ''), 6: bool(c_act)
                })
        conexion.close()
    except Exception as e:
        logger.error(f"Error cargando clientes: {e}")
        
    return render_template('clientes.html', clientes=clientes_lista)


@app.route('/clientes/crear', methods=['POST'], endpoint='crear_cliente')
@app.route('/clientes/agregar', methods=['POST'], endpoint='agregar_cliente')
def crear_cliente():
    if session.get('rol') not in ['admin', 'empleado', 'dueño']:
        if request.is_json or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'success': False, 'error': 'No autorizado'}), 403
        flash('No autorizado.', 'error')
        return redirect(url_for('clientes'))

    tenant_id = session.get('tenant_id', 1)

    try:
        req_json = request.get_json(silent=True) or {}
        
        def get_field(*keys, default=''):
            for k in keys:
                v = req_json.get(k)
                if v is not None and str(v).strip() != '':
                    return str(v).strip()
                v = request.form.get(k)
                if v is not None and str(v).strip() != '':
                    return str(v).strip()
            return default

        nombre = get_field('nombre', 'cliente_nombre')
        email = get_field('email', 'correo')
        telefono = get_field('telefono', 'celular')
        direccion = get_field('direccion')
        documento = get_field('documento', 'dni', 'ruc')

        if not nombre:
            msg = 'El nombre del cliente es obligatorio.'
            if request.is_json or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return jsonify({'success': False, 'error': msg}), 400
            flash(msg, 'warning')
            return redirect(url_for('clientes'))

        if not email:
            clean_n = ''.join(e for e in nombre if e.isalnum()).lower()
            email = f"{clean_n}_{int(time.time())}_{random.randint(100,999)}@cliente.local"

        conexion = obtener_conexion()
        with conexion.cursor() as cursor:
            cursor.execute("""
                INSERT INTO clientes (nombre, email, telefono, direccion, documento, activo, tenant_id)
                VALUES (%s, %s, %s, %s, %s, true, %s)
                RETURNING id
            """, (nombre, email, telefono, direccion, documento, tenant_id))
            nuevo_id = cursor.fetchone()[0]

        conexion.commit()
        conexion.close()

        msg = f'Cliente "{nombre}" registrado exitosamente.'
        if request.is_json or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'success': True, 'message': msg, 'cliente_id': nuevo_id})

        flash(msg, 'success')
    except Exception as e:
        logger.error(f"Error al registrar cliente: {e}")
        if request.is_json or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'success': False, 'error': str(e)}), 500
        flash(f'Error al registrar cliente: {e}', 'error')

    return redirect(url_for('clientes'))


@app.route('/clientes/editar', methods=['POST'], endpoint='editar_cliente')
def editar_cliente():
    if session.get('rol') not in ['admin', 'empleado', 'dueño']:
        if request.is_json or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'success': False, 'error': 'No autorizado'}), 403
        flash('No autorizado.', 'error')
        return redirect(url_for('clientes'))

    tenant_id = session.get('tenant_id', 1)

    try:
        req_json = request.get_json(silent=True) or {}
        
        def get_field(*keys, default=''):
            for k in keys:
                v = req_json.get(k)
                if v is not None and str(v).strip() != '':
                    return str(v).strip()
                v = request.form.get(k)
                if v is not None and str(v).strip() != '':
                    return str(v).strip()
            return default

        raw_id = get_field('id', 'cliente_id')
        cliente_id = int(raw_id) if raw_id and str(raw_id).isdigit() else None

        nombre = get_field('nombre', 'cliente_nombre')
        email = get_field('email', 'correo')
        telefono = get_field('telefono', 'celular')
        direccion = get_field('direccion')
        documento = get_field('documento', 'dni', 'ruc')

        if not cliente_id or not nombre:
            msg = 'El nombre y el ID de cliente son obligatorios.'
            if request.is_json or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return jsonify({'success': False, 'error': msg}), 400
            flash(msg, 'warning')
            return redirect(url_for('clientes'))

        conexion = obtener_conexion()
        with conexion.cursor() as cursor:
            cursor.execute("""
                UPDATE clientes
                SET nombre = %s, email = COALESCE(NULLIF(%s, ''), email), telefono = %s, direccion = %s, documento = %s
                WHERE id = %s AND (tenant_id = %s OR tenant_id IS NULL)
            """, (nombre, email, telefono, direccion, documento, cliente_id, tenant_id))

        conexion.commit()
        conexion.close()

        msg = 'Cliente actualizado exitosamente.'
        if request.is_json or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'success': True, 'message': msg})

        flash(msg, 'success')
    except Exception as e:
        logger.error(f"Error al editar cliente: {e}")
        if request.is_json or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'success': False, 'error': str(e)}), 500
        flash(f'Error al editar cliente: {e}', 'error')

    return redirect(url_for('clientes'))


@app.route('/clientes/detalle', defaults={'cliente_id': None})
@app.route('/clientes/detalle/<int:cliente_id>')
@app.route('/clientes/<int:cliente_id>')
def detalle_cliente(cliente_id):
    if 'rol' not in session or session['rol'] not in ['admin', 'empleado', 'dueño']:
        if request.is_json or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'success': False, 'error': 'Acceso denegado.'}), 403
        flash('Acceso denegado.', 'error')
        return redirect(url_for('dashboard'))

    if cliente_id is None:
        raw_id = request.args.get('cliente_id') or request.args.get('id')
        cliente_id = int(raw_id) if raw_id and str(raw_id).isdigit() else None

    if not cliente_id:
        if request.is_json or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'success': False, 'error': 'ID no especificado'}), 400
        flash('No se especificó un cliente válido.', 'warning')
        return redirect(url_for('clientes'))

    tenant_id = session.get('tenant_id', 1)
    cliente = None
    mascotas_lista = []

    try:
        conexion = obtener_conexion()
        with conexion.cursor() as cursor:
            cursor.execute("""
                SELECT id, COALESCE(nombre, '') AS nombre, COALESCE(email, '') AS email, 
                       COALESCE(telefono, '') AS telefono, COALESCE(direccion, '') AS direccion, 
                       COALESCE(documento, '') AS documento, COALESCE(activo, true) AS activo
                FROM clientes
                WHERE id = %s AND (tenant_id = %s OR tenant_id IS NULL)
            """, (cliente_id, tenant_id))
            
            r = cursor.fetchone()
            if r:
                cliente = {
                    'id': r[0], 'nombre': r[1], 'email': r[2], 'telefono': r[3],
                    'direccion': r[4], 'documento': r[5], 'activo': r[6],
                    0: r[0], 1: r[1], 2: r[2], 3: r[3], 4: r[4], 5: r[5]
                }

            cursor.execute("""
                SELECT id, nombre, COALESCE(especie, '') AS especie, COALESCE(raza, '') AS raza
                FROM mascotas
                WHERE cliente_id = %s AND (tenant_id = %s OR tenant_id IS NULL)
            """, (cliente_id, tenant_id))
            for m in cursor.fetchall():
                mascotas_lista.append({'id': m[0], 'nombre': m[1], 'especie': m[2], 'raza': m[3]})

        conexion.close()
    except Exception as e:
        logger.error(f"Error consultando detalle cliente {cliente_id}: {e}")

    if not cliente:
        if request.is_json or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'success': False, 'error': 'Cliente no encontrado'}), 404
        flash('Cliente no encontrado.', 'error')
        return redirect(url_for('clientes'))

    if request.is_json or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
        return jsonify({'success': True, 'cliente': cliente, 'mascotas': mascotas_lista})

    return render_template('detalle_cliente.html', cliente=cliente, mascotas=mascotas_lista)


# ─── GESTIÓN DE PERSONAL Y USUARIOS ──────────────────────────────────────────

@app.route('/personal', methods=['GET', 'POST'])
@app.route('/personal/agregar', methods=['POST'], endpoint='agregar_personal')
@app.route('/personal/editar', methods=['POST'], endpoint='editar_personal')
def personal():
    if 'rol' not in session or session['rol'] not in ['admin', 'dueño']:
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest' or request.is_json:
            return jsonify({'success': False, 'error': 'Acceso denegado.'}), 403
        flash('Acceso denegado.', 'error')
        return redirect(url_for('dashboard'))

    tenant_id = session.get('tenant_id', 1)

    if request.method == 'POST':
        try:
            req_json = request.get_json(silent=True) or {}
            
            def get_param(*keys, default=''):
                for k in keys:
                    v = req_json.get(k)
                    if v is not None and str(v).strip() != '':
                        return str(v).strip()
                    v = request.form.get(k)
                    if v is not None and str(v).strip() != '':
                        return str(v).strip()
                return default

            raw_id = get_param('id', 'personal_id', 'usuario_id', 'edit_id')
            target_id = int(raw_id) if raw_id and raw_id.isdigit() else None

            if 'eliminar' in request.form or get_param('accion') in ['desactivar', 'eliminar']:
                if target_id:
                    conexion = obtener_conexion()
                    with conexion.cursor() as cursor:
                        cursor.execute("""
                            UPDATE personal 
                            SET activo = false 
                            WHERE (id = %s OR usuario_id = %s) AND (tenant_id = %s OR tenant_id IS NULL)
                            RETURNING usuario_id
                        """, (target_id, target_id, tenant_id))
                        res = cursor.fetchone()
                        if res and res[0]:
                            cursor.execute("UPDATE usuarios SET activo = false WHERE id = %s", (res[0],))
                    conexion.commit()
                    conexion.close()
                    flash('Empleado desactivado correctamente.', 'success')
                    return redirect(url_for('personal'))

            username = get_param('username', 'nombre', 'usuario')
            cargo = get_param('cargo', default='Empleado')
            rol = get_param('rol', default='empleado').lower()
            rol_final = rol if rol in ['cliente', 'dueño', 'admin', 'empleado', 'superadmin'] else 'empleado'

            salario_raw = get_param('salario', 'sueldo')
            salario = float(salario_raw) if salario_raw else 0.0

            conexion = obtener_conexion()
            with conexion.cursor() as cursor:
                if target_id:
                    cursor.execute("""
                        SELECT id, usuario_id 
                        FROM personal 
                        WHERE (id = %s OR usuario_id = %s) AND (tenant_id = %s OR tenant_id IS NULL)
                    """, (target_id, target_id, tenant_id))
                    p_row = cursor.fetchone()
                    
                    if p_row:
                        p_id, usuario_id = p_row[0], p_row[1]
                        cursor.execute("""
                            UPDATE personal SET cargo = %s, salario = %s 
                            WHERE id = %s AND (tenant_id = %s OR tenant_id IS NULL)
                        """, (cargo, salario, p_id, tenant_id))
                    else:
                        usuario_id = target_id
                        cursor.execute("""
                            INSERT INTO personal (usuario_id, cargo, salario, activo, tenant_id) 
                            VALUES (%s, %s, %s, true, %s)
                        """, (usuario_id, cargo, salario, tenant_id))

                    if usuario_id and username:
                        cursor.execute("""
                            UPDATE usuarios SET username = %s, rol = %s::rol_usuario_enum 
                            WHERE id = %s AND (tenant_id = %s OR tenant_id IS NULL)
                        """, (username, rol_final, usuario_id, tenant_id))

                    mensaje = 'Datos del empleado actualizados correctamente.'
                else:
                    if not username:
                        conexion.close()
                        msg = 'El nombre de usuario es obligatorio.'
                        if request.is_json:
                            return jsonify({'success': False, 'error': msg}), 400
                        flash(msg, 'warning')
                        return redirect(url_for('personal'))

                    from controladores.usuarios_controlador import hash_password
                    password_default = hash_password('123456')

                    cursor.execute("""
                        INSERT INTO usuarios (username, password, rol, activo, tenant_id)
                        VALUES (%s, %s, %s::rol_usuario_enum, true, %s)
                        RETURNING id
                    """, (username, password_default, rol_final, tenant_id))
                    nuevo_usuario_id = cursor.fetchone()[0]

                    cursor.execute("""
                        INSERT INTO personal (usuario_id, cargo, salario, activo, tenant_id)
                        VALUES (%s, %s, %s, true, %s)
                    """, (nuevo_usuario_id, cargo, salario, tenant_id))

                    mensaje = f'Empleado "{username}" registrado correctamente.'

            conexion.commit()
            conexion.close()

            if request.headers.get('X-Requested-With') == 'XMLHttpRequest' or request.is_json:
                return jsonify({'success': True, 'message': mensaje})

            flash(mensaje, 'success')
        except Exception as e:
            logger.error(f"Error procesando personal: {e}")
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest' or request.is_json:
                return jsonify({'success': False, 'error': str(e)}), 500
            flash(f'Error al procesar personal: {e}', 'error')

        return redirect(url_for('personal'))

    empleados_lista = []
    usuarios_lista = []

    try:
        conexion = obtener_conexion()
        with conexion.cursor() as cursor:
            cursor.execute("""
                SELECT p.id AS personal_id, u.username, COALESCE(p.cargo, 'Empleado') AS cargo,
                       u.rol, COALESCE(p.salario, 0.00) AS salario, COALESCE(p.activo, true) AS activo, u.id AS usuario_id
                FROM personal p
                INNER JOIN usuarios u ON p.usuario_id = u.id
                WHERE (p.tenant_id = %s OR p.tenant_id IS NULL)
                ORDER BY p.id ASC
            """, (tenant_id,))
            
            for r in cursor.fetchall():
                p_id, username, cargo, rol, salario, activo_bool, usuario_id = r[0], r[1] or 'Sin Usuario', r[2] or 'Empleado', str(r[3]), float(r[4] or 0), bool(r[5]), r[6]
                estado_str = 'Activo' if activo_bool else 'Inactivo'

                empleados_lista.append({
                    'id': p_id, 'personal_id': p_id, 'usuario': username, 'nombre': username, 'username': username,
                    'cargo': cargo, 'rol': rol, 'salario': salario, 'sueldo': salario, 'activo': 1 if activo_bool else 0,
                    'estado': estado_str, 'usuario_id': usuario_id,
                    0: p_id, 1: usuario_id, 2: username, 3: cargo, 4: salario, 5: 1 if activo_bool else 0, 6: rol, 7: estado_str
                })

            cursor.execute("""
                SELECT id, username, rol, COALESCE(activo, true) 
                FROM usuarios 
                WHERE (tenant_id = %s OR tenant_id IS NULL)
                ORDER BY id ASC
            """, (tenant_id,))
            for u in cursor.fetchall():
                usuarios_lista.append({'id': u[0], 'username': u[1], 'rol': str(u[2]), 'activo': u[3], 0: u[0], 1: u[1], 2: str(u[2]), 3: u[3]})

        conexion.close()
    except Exception as e:
        logger.error(f"Error cargando módulo personal: {e}")

    puede_modificar = session.get('rol') in ['admin', 'dueño']

    return render_template('personal.html', empleados=empleados_lista, personal=empleados_lista, usuarios=usuarios_lista, puede_modificar=puede_modificar)


@app.route('/personal/toggle-estado/<int:target_id>', methods=['POST'])
@app.route('/usuario/toggle-estado/<int:target_id>', methods=['POST'])
def toggle_estado_personal(target_id):
    if session.get('rol') not in ['admin', 'dueño']:
        return jsonify({'success': False, 'error': 'Sin permisos'}), 403

    tenant_id = session.get('tenant_id', 1)

    try:
        conexion = obtener_conexion()
        with conexion.cursor() as cursor:
            cursor.execute("""
                UPDATE personal SET activo = NOT COALESCE(activo, true)
                WHERE (id = %s OR usuario_id = %s) AND (tenant_id = %s OR tenant_id IS NULL)
                RETURNING usuario_id, activo
            """, (target_id, target_id, tenant_id))
            
            res = cursor.fetchone()
            if res and res[0]:
                cursor.execute("UPDATE usuarios SET activo = %s WHERE id = %s", (res[1], res[0]))
                nuevo_estado = res[1]
            else:
                cursor.execute("UPDATE usuarios SET activo = NOT COALESCE(activo, true) WHERE id = %s RETURNING activo", (target_id,))
                res_u = cursor.fetchone()
                if not res_u:
                    conexion.close()
                    return jsonify({'success': False, 'error': 'Usuario no encontrado'}), 404
                nuevo_estado = res_u[0]

        conexion.commit()
        conexion.close()

        if request.is_json or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'success': True, 'nuevo_estado': nuevo_estado, 'message': 'Estado de acceso actualizado.'})

        flash('Estado de acceso actualizado correctamente.', 'success')
        return redirect(url_for('personal'))
    except Exception as e:
        logger.error(f"Error al cambiar estado de acceso: {e}")
        if request.is_json or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'success': False, 'error': str(e)}), 500
        flash(f'Error al cambiar estado: {e}', 'error')
        return redirect(url_for('personal'))


@app.route('/personal/eliminar/<int:target_id>', methods=['POST', 'DELETE'])
@app.route('/usuario/eliminar/<int:target_id>', methods=['POST', 'DELETE'])
def eliminar_personal_permanente(target_id):
    if session.get('rol') not in ['admin', 'dueño', 'superadmin']:
        if request.is_json or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'success': False, 'error': 'Sin permisos'}), 403
        flash('Sin permisos para eliminar personal.', 'error')
        return redirect(url_for('personal'))

    try:
        conexion = obtener_conexion()
        with conexion.cursor() as cursor:
            cursor.execute("SELECT p.id AS personal_id, u.id AS usuario_id FROM usuarios u LEFT JOIN personal p ON p.usuario_id = u.id WHERE u.id = %s OR p.id = %s", (target_id, target_id))
            row = cursor.fetchone()
            personal_id = row[0] if row else None
            usuario_id = row[1] if row else target_id

            if personal_id:
                cursor.execute("DELETE FROM asistencia WHERE personal_id = %s", (personal_id,))

            if usuario_id:
                cursor.execute("DELETE FROM asistencia WHERE personal_id IN (SELECT id FROM personal WHERE usuario_id = %s)", (usuario_id,))
                cursor.execute("DELETE FROM ventas WHERE vendedor_id = %s", (usuario_id,))
                cursor.execute("DELETE FROM personal WHERE usuario_id = %s", (usuario_id,))
                cursor.execute("DELETE FROM usuarios WHERE id = %s", (usuario_id,))

        conexion.commit()
        conexion.close()

        if request.is_json or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'success': True, 'message': 'Usuario y su historial eliminados.'})

        flash('Usuario eliminado permanentemente.', 'success')
    except Exception as e:
        logger.error(f"Error eliminando personal: {e}")
        if request.is_json or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'success': False, 'error': str(e)}), 500
        flash(f'Error al eliminar: {e}', 'error')

    return redirect(url_for('personal'))


# ─── ASISTENCIA DEL PERSONAL CON ZONA HORARIA DE PERÚ ────────────────────────

def obtener_usuario_id_sesion():
    uid = session.get('usuario_id') or session.get('user_id') or session.get('id')
    if uid and str(uid).isdigit():
        return int(uid)
    
    username = session.get('usuario') or session.get('username')
    if username:
        try:
            conexion = obtener_conexion()
            with conexion.cursor() as cursor:
                cursor.execute("SELECT id FROM usuarios WHERE username = %s OR id::text = %s LIMIT 1", (str(username), str(username)))
                row = cursor.fetchone()
                if row:
                    conexion.close()
                    return row[0]
            conexion.close()
        except Exception as e:
            logger.error(f"Error resolviendo ID de usuario: {e}")
            
    return None


@app.route('/asistencia', methods=['GET', 'POST'])
def asistencia():
    if 'rol' not in session:
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest' or request.is_json:
            return jsonify({'success': False, 'error': 'No autorizado'}), 401
        flash('Debes iniciar sesión para acceder a la asistencia.', 'error')
        return redirect(url_for('login'))
        
    tenant_id = session.get('tenant_id', 1)
    usuario_id = obtener_usuario_id_sesion()
    rol = session.get('rol', 'empleado')
    
    hoy_peru = datetime.now(ZONA_HORARIA_PERU).date()

    if not usuario_id:
        flash('Sesión de usuario no válida.', 'error')
        return redirect(url_for('login'))

    if request.method == 'POST':
        return _procesar_marcar_asistencia(usuario_id, tenant_id)

    registros = []
    asistencia_actual = None

    try:
        conexion = obtener_conexion()
        with conexion.cursor() as cursor:
            cursor.execute("""
                SELECT a.id, TO_CHAR(a.hora_entrada, 'HH12:MI AM') AS entrada_fmt, TO_CHAR(a.hora_salida, 'HH12:MI AM') AS salida_fmt
                FROM asistencia a
                INNER JOIN personal p ON a.personal_id = p.id
                WHERE p.usuario_id = %s AND a.fecha = %s AND a.hora_salida IS NULL
                ORDER BY a.id DESC LIMIT 1
            """, (usuario_id, hoy_peru))
            a_row = cursor.fetchone()
            
            if a_row:
                asistencia_actual = {'id': a_row[0], 'hora_entrada': a_row[1], 'hora_salida': a_row[2], 'estado': 'presente'}

            if rol in ['admin', 'dueño']:
                cursor.execute("""
                    SELECT a.id, u.username, COALESCE(p.cargo, 'Empleado') AS cargo, a.fecha,
                           TO_CHAR(a.hora_entrada, 'HH12:MI AM') AS entrada, TO_CHAR(a.hora_salida, 'HH12:MI AM') AS salida,
                           CASE WHEN a.hora_salida IS NOT NULL THEN 'completado' ELSE 'presente' END AS estado
                    FROM asistencia a
                    INNER JOIN personal p ON a.personal_id = p.id
                    INNER JOIN usuarios u ON p.usuario_id = u.id
                    WHERE (a.tenant_id = %s OR a.tenant_id IS NULL) AND a.fecha = %s
                    ORDER BY a.hora_entrada DESC
                """, (tenant_id, hoy_peru))
            else:
                cursor.execute("""
                    SELECT a.id, u.username, COALESCE(p.cargo, 'Empleado') AS cargo, a.fecha,
                           TO_CHAR(a.hora_entrada, 'HH12:MI AM') AS entrada, TO_CHAR(a.hora_salida, 'HH12:MI AM') AS salida,
                           CASE WHEN a.hora_salida IS NOT NULL THEN 'completado' ELSE 'presente' END AS estado
                    FROM asistencia a
                    INNER JOIN personal p ON a.personal_id = p.id
                    INNER JOIN usuarios u ON p.usuario_id = u.id
                    WHERE (a.tenant_id = %s OR a.tenant_id IS NULL) AND p.usuario_id = %s AND a.fecha = %s
                    ORDER BY a.hora_entrada DESC
                """, (tenant_id, usuario_id, hoy_peru))

            for r in cursor.fetchall():
                registros.append({
                    'id': r[0], 'usuario': r[1] or 'Empleado', 'nombre': r[1] or 'Empleado', 'username': r[1] or 'Empleado',
                    'cargo': r[2] or 'Empleado', 'fecha': str(r[3]), 'hora_entrada': r[4] or '--:--', 'entrada': r[4] or '--:--',
                    'hora_salida': r[5] or 'En turno', 'salida': r[5] or 'En turno', 'estado': r[6],
                    0: r[0], 1: r[1], 2: r[2], 3: str(r[3]), 4: r[4] or '--:--', 5: r[5] or 'En turno', 6: r[6]
                })

        conexion.close()
    except Exception as e:
        logger.error(f"Error cargando asistencia: {e}")

    return render_template('asistencia.html', registros=registros, historial=registros, asistencia=asistencia_actual, asistencia_actual=asistencia_actual)


@app.route('/asistencia/marcar', methods=['GET', 'POST'])
@app.route('/marcar_asistencia', methods=['GET', 'POST'])
def marcar_asistencia():
    if 'rol' not in session:
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest' or request.is_json:
            return jsonify({'success': False, 'error': 'No autorizado'}), 401
        flash('Debes iniciar sesión para registrar asistencia.', 'error')
        return redirect(url_for('login'))

    if request.method == 'GET':
        return redirect(url_for('asistencia'))

    tenant_id = session.get('tenant_id', 1)
    usuario_id = obtener_usuario_id_sesion()
    
    if not usuario_id:
        msg = 'Sesión no válida.'
        return jsonify({'success': False, 'error': msg}) if request.is_json else (flash(msg, 'error') or redirect(url_for('login')))

    return _procesar_marcar_asistencia(usuario_id, tenant_id)


def _procesar_marcar_asistencia(usuario_id, tenant_id):
    try:
        ahora_peru = datetime.now(ZONA_HORARIA_PERU)
        fecha_hoy = ahora_peru.date()
        hora_actual = ahora_peru.time()

        conexion = obtener_conexion()
        with conexion.cursor() as cursor:
            cursor.execute("SELECT id FROM personal WHERE usuario_id = %s ORDER BY id DESC LIMIT 1", (usuario_id,))
            p_row = cursor.fetchone()

            if p_row:
                personal_id = p_row[0]
            else:
                cursor.execute("""
                    INSERT INTO personal (usuario_id, cargo, salario, activo, tenant_id)
                    VALUES (%s, 'Empleado', 0.00, true, %s)
                    RETURNING id
                """, (usuario_id, tenant_id))
                personal_id = cursor.fetchone()[0]

            cursor.execute("SELECT id, hora_salida FROM asistencia WHERE personal_id = %s AND fecha = %s LIMIT 1", (personal_id, fecha_hoy))
            reg_hoy = cursor.fetchone()

            if not reg_hoy:
                cursor.execute("INSERT INTO asistencia (personal_id, fecha, hora_entrada, tenant_id) VALUES (%s, %s, %s, %s)", (personal_id, fecha_hoy, hora_actual, tenant_id))
                mensaje = 'Hora de entrada registrada correctamente.'
            else:
                asistencia_id, hora_salida_existente = reg_hoy[0], reg_hoy[1]
                if hora_salida_existente is None:
                    cursor.execute("UPDATE asistencia SET hora_salida = %s WHERE id = %s", (hora_actual, asistencia_id))
                    mensaje = 'Hora de salida registrada correctamente.'
                else:
                    cursor.execute("UPDATE asistencia SET hora_entrada = %s, hora_salida = NULL WHERE id = %s", (hora_actual, asistencia_id))
                    mensaje = 'Nueva hora de entrada registrada correctamente.'

        conexion.commit()
        conexion.close()

        if request.headers.get('X-Requested-With') == 'XMLHttpRequest' or request.is_json:
            return jsonify({'success': True, 'message': mensaje})
            
        flash(mensaje, 'success')
    except Exception as e:
        logger.error(f"Error registrando asistencia: {e}")
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest' or request.is_json:
            return jsonify({'success': False, 'error': str(e)}), 500
        flash('Error al registrar asistencia.', 'error')

    return redirect(url_for('asistencia'))


# ─── GESTIÓN DE CITAS ────────────────────────────────────────────────────────

@app.route('/citas')
def citas():
    if 'rol' not in session or session['rol'] not in ['admin', 'empleado', 'dueño']:
        flash('Acceso denegado.', 'error')
        return redirect(url_for('dashboard'))
    
    tenant_id = session.get('tenant_id', 1)
    hoy_str = datetime.now(ZONA_HORARIA_PERU).strftime('%Y-%m-%d')
    fecha_filtro = request.args.get('fecha', default=hoy_str)

    citas_list = []
    servicios_lista = []

    try:
        conexion = obtener_conexion()
        with conexion.cursor() as cursor:
            cursor.execute("""
                SELECT c.id, c.cliente_nombre, COALESCE(c.cliente_email, c.cliente_telefono), c.cliente_telefono,
                       TO_CHAR(c.hora, 'HH12:MI AM') AS hora_fmt, c.estado, COALESCE(s.nombre, 'Servicio General') AS s_nom,
                       c.mascota_nombre, c.mascota_especie, COALESCE(c.precio_total, 0.00) AS precio, c.fecha, c.observaciones
                FROM citas c
                LEFT JOIN servicios s ON c.servicio_id = s.id
                WHERE c.tenant_id = %s AND c.fecha = %s
                ORDER BY c.hora ASC
            """, (tenant_id, fecha_filtro))
            
            for r in cursor.fetchall():
                citas_list.append({
                    'id': r[0], 'cliente_nombre': r[1] or '', 'cliente_email': r[2] or '', 'cliente_telefono': r[3] or '',
                    'hora': r[4] or '', 'estado': str(r[5]).lower() if r[5] else 'pendiente', 'servicio_nombre': r[6],
                    'mascota_nombre': r[7] or '', 'mascota_especie': r[8] or '', 'precio_total': float(r[9] or 0),
                    'fecha': str(r[10]) if r[10] else '', 'observaciones': r[11] or '',
                    0: r[0], 1: r[1] or '', 2: r[2] or '', 3: r[3] or '', 4: r[4] or '', 5: str(r[5]).lower() if r[5] else 'pendiente',
                    6: r[6], 7: r[7] or '', 8: r[8] or '', 9: float(r[9] or 0), 10: str(r[10]) if r[10] else '', 11: r[11] or ''
                })

            cursor.execute("SELECT id, nombre, precio, duracion, max_citas_dia, COALESCE(activo, true) FROM servicios WHERE tenant_id = %s ORDER BY nombre ASC", (tenant_id,))
            for s in cursor.fetchall():
                servicios_lista.append({'id': s[0], 'nombre': s[1], 'precio': float(s[2]) if s[2] else 0.0, 0: s[0], 1: s[1], 2: float(s[2]) if s[2] else 0.0, 5: bool(s[5])})

        conexion.close()
    except Exception as e:
        logger.error(f"Error consultando citas: {e}")

    return render_template('citas.html', citas=citas_list, fecha_filtro=fecha_filtro, servicios=servicios_lista)


@app.route('/agendar_cita', methods=['POST'])
def agendar_cita():
    try:
        tenant_id = session.get('tenant_id', 1)
        nombre = request.form.get('cliente_nombre', '').strip()
        email = request.form.get('cliente_email', '').strip() or None
        telefono = request.form.get('cliente_telefono', '').strip() or None
        direccion = request.form.get('cliente_direccion', '').strip() or None
        
        fecha = request.form.get('fecha', '').strip()
        hora = request.form.get('hora', '').strip()
        
        servicio_id_raw = request.form.get('servicio_id') or request.form.get('servicio')
        if not servicio_id_raw and request.form.getlist('servicio_id[]'):
            servicio_id_raw = request.form.getlist('servicio_id[]')[0]
            
        servicio_id = int(servicio_id_raw) if servicio_id_raw and str(servicio_id_raw).isdigit() else None
        mascota_nombre = request.form.get('mascota_nombre', '').strip() or None
        mascota_especie = request.form.get('mascota_especie', 'perro').strip() or 'perro'
        
        obs_input = request.form.get('observaciones', '').strip()
        obs_partes = []
        if direccion:
            obs_partes.append(f"Dirección: {direccion}")
        if obs_input:
            obs_partes.append(obs_input)
        observaciones = " | ".join(obs_partes) if obs_partes else None

        if not nombre or not fecha or not hora or not servicio_id or not mascota_nombre:
            flash('Por favor completa todos los campos obligatorios.', 'error')
            return redirect(request.referrer or url_for('citas'))

        conexion = obtener_conexion()
        with conexion.cursor() as cursor:
            precio_total = 0.0
            cursor.execute("SELECT precio FROM servicios WHERE id = %s", (servicio_id,))
            s_row = cursor.fetchone()
            if s_row and s_row[0] is not None:
                precio_total = float(s_row[0])

            cursor.execute("""
                INSERT INTO citas (
                    cliente_nombre, cliente_email, cliente_telefono,
                    servicio_id, fecha, hora, mascota_nombre, mascota_especie,
                    observaciones, precio_total, estado, tenant_id
                ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, 'pendiente'::estado_cita_enum, %s)
            """, (nombre, email, telefono, servicio_id, fecha, hora, mascota_nombre, mascota_especie, observaciones, precio_total, tenant_id))

        conexion.commit()
        conexion.close()

        flash('Cita agendada exitosamente.', 'success')
    except Exception as e:
        logger.error(f"Error al agendar cita: {e}")
        flash(f'Error al agendar cita: {str(e)}', 'error')

    return redirect(request.referrer or url_for('citas'))


@app.route('/cita/<int:cita_id>/status', methods=['POST'])
@app.route('/citas/estado/<int:cita_id>', methods=['POST'])
def cambiar_estado_cita(cita_id):
    if session.get('rol') not in ['admin', 'empleado', 'dueño']:
        return jsonify({'success': False, 'error': 'Sin permisos'}), 403

    try:
        tenant_id = session.get('tenant_id', 1)
        data = request.get_json(silent=True) or {}
        nuevo_estado = (data.get('status') or data.get('estado') or request.form.get('estado') or '').strip().lower()
        motivo = (data.get('motivo') or request.form.get('motivo') or '').strip()

        if nuevo_estado not in ['pendiente', 'confirmada', 'en_progreso', 'completada', 'cancelada']:
            return jsonify({'success': False, 'error': 'Estado inválido'}), 400

        conexion = obtener_conexion()
        with conexion.cursor() as cursor:
            cursor.execute("UPDATE citas SET estado = %s::estado_cita_enum WHERE id = %s AND tenant_id = %s", (nuevo_estado, cita_id, tenant_id))
            if motivo:
                cursor.execute("UPDATE citas SET observaciones = COALESCE(observaciones, '') || %s WHERE id = %s AND tenant_id = %s", (f" | Motivo: {motivo}", cita_id, tenant_id))

        conexion.commit()
        conexion.close()

        return jsonify({'success': True, 'message': f'Estado actualizado a {nuevo_estado}.'})
    except Exception as e:
        logger.error(f"Error cambiando estado de cita: {e}")
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/cita/<int:cita_id>/reprogramar', methods=['POST'])
def reprogramar_cita(cita_id):
    if session.get('rol') not in ['admin', 'empleado', 'dueño']:
        return jsonify({'success': False, 'error': 'Sin permisos'}), 403

    tenant_id = session.get('tenant_id', 1)
    data = request.get_json(silent=True) or {}
    nueva_fecha = (data.get('fecha') or request.form.get('fecha') or '').strip()
    nueva_hora = (data.get('hora') or request.form.get('hora') or '').strip()

    if not nueva_fecha or not nueva_hora:
        return jsonify({'success': False, 'error': 'Fecha y hora requeridas.'}), 400

    conexion = obtener_conexion()
    try:
        with conexion.cursor() as cursor:
            cursor.execute(
                "UPDATE citas SET fecha = %s, hora = %s WHERE id = %s AND tenant_id = %s",
                (nueva_fecha, nueva_hora, cita_id, tenant_id)
            )
        conexion.commit()
        return jsonify({'success': True, 'message': 'Cita reprogramada correctamente.'})
    except Exception as e:
        conexion.rollback()
        logger.error(f"Error reprogramando cita ID={cita_id}: {e}")
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conexion.close()


@app.route('/cita/<int:cita_id>', endpoint='ver_detalles_cita')
@app.route('/cita/<int:cita_id>/recibo', endpoint='ver_recibo_cita')
@app.route('/cita/recibo/<int:cita_id>')
@app.route('/cita/<int:cita_id>/detalles')
def ver_recibo_cita(cita_id):
    if 'rol' not in session:
        return redirect(url_for('login'))
    
    tenant_id = session.get('tenant_id', 1)
    cita = None

    conexion = obtener_conexion()
    try:
        with conexion.cursor() as cursor:
            cursor.execute("""
                SELECT c.id, c.cliente_nombre, COALESCE(c.cliente_email, ''), COALESCE(c.cliente_telefono, ''),
                       TO_CHAR(c.fecha, 'YYYY-MM-DD'), TO_CHAR(c.hora, 'HH12:MI AM'), COALESCE(c.mascota_nombre, ''),
                       COALESCE(c.mascota_especie, 'perro'), COALESCE(c.observaciones, ''), COALESCE(c.precio_total, 0.00),
                       COALESCE(c.estado, 'pendiente'), COALESCE(s.nombre, 'Servicio General'), COALESCE(s.precio, 0.00)
                FROM citas c
                LEFT JOIN servicios s ON c.servicio_id = s.id
                WHERE c.id = %s AND c.tenant_id = %s
            """, (cita_id, tenant_id))
            
            r = cursor.fetchone()
            if r:
                cita = {
                    'id': r[0], 'cliente_nombre': r[1] or 'Cliente General', 'cliente_email': r[2], 'cliente_telefono': r[3],
                    'fecha': r[4] or '', 'hora': r[5] or '', 'mascota_nombre': r[6], 'mascota_especie': r[7],
                    'observaciones': r[8], 'precio_total': float(r[9] or 0), 'estado': str(r[10]).lower(),
                    'servicio_nombre': r[11], 'servicio_precio': float(r[12] or 0),
                    0: r[0], 1: r[1], 2: r[2], 3: r[3], 4: r[4], 5: r[5], 6: r[6], 7: r[7], 8: r[8], 9: float(r[9] or 0), 10: str(r[10]).lower()
                }
    except Exception as e:
        logger.error(f"Error recibo cita ID={cita_id}: {e}")
    finally:
        conexion.close()

    if not cita:
        flash('Cita no encontrada.', 'error')
        return redirect(url_for('citas'))

    return render_template('recibo_cita.html', cita=cita, momento_actual=datetime.now(ZONA_HORARIA_PERU))


@app.route('/cita/ticket/<int:cita_id>', endpoint='ticket_cita')
@app.route('/cita/<int:cita_id>/ticket')
def ver_ticket_cita(cita_id):
    if 'rol' not in session:
        return redirect(url_for('login'))
    
    tenant_id = session.get('tenant_id', 1)
    cita = None

    conexion = obtener_conexion()
    try:
        with conexion.cursor() as cursor:
            cursor.execute("""
                SELECT c.id, c.cliente_nombre, COALESCE(c.cliente_email, ''), COALESCE(c.cliente_telefono, ''),
                       TO_CHAR(c.fecha, 'YYYY-MM-DD'), TO_CHAR(c.hora, 'HH12:MI AM'), COALESCE(c.mascota_nombre, ''),
                       COALESCE(c.mascota_especie, 'perro'), COALESCE(c.observaciones, ''), COALESCE(c.precio_total, 0.00),
                       COALESCE(c.estado, 'pendiente'), COALESCE(s.nombre, 'Servicio General')
                FROM citas c
                LEFT JOIN servicios s ON c.servicio_id = s.id
                WHERE c.id = %s AND c.tenant_id = %s
            """, (cita_id, tenant_id))
            r = cursor.fetchone()
            if r:
                cita = {
                    'id': r[0], 'cliente_nombre': r[1] or 'Cliente General', 'cliente_email': r[2], 'cliente_telefono': r[3],
                    'fecha': r[4] or '', 'hora': r[5] or '', 'mascota_nombre': r[6], 'mascota_especie': r[7],
                    'observaciones': r[8], 'precio_total': float(r[9] or 0), 'estado': str(r[10]).lower(), 'servicio_nombre': r[11]
                }
    except Exception as e:
        logger.error(f"Error ticket cita ID={cita_id}: {e}")
    finally:
        conexion.close()

    if not cita:
        flash('Cita no encontrada.', 'error')
        return redirect(url_for('citas'))

    return render_template('ticket_cita.html', cita=cita, momento_actual=datetime.now(ZONA_HORARIA_PERU))


@app.route('/cita/<int:cita_id>/editar', methods=['GET', 'POST'])
def editar_cita(cita_id):
    if session.get('rol') not in ['admin', 'empleado', 'dueño']:
        flash('Sin permisos.', 'error')
        return redirect(url_for('citas'))

    tenant_id = session.get('tenant_id', 1)

    if request.method == 'POST':
        conexion = obtener_conexion()
        try:
            nombre = request.form.get('cliente_nombre', '').strip()
            fecha = request.form.get('fecha', '').strip()
            hora = request.form.get('hora', '').strip()
            servicio_id = request.form.get('servicio_id')
            mascota_nombre = request.form.get('mascota_nombre', '').strip()
            observaciones = request.form.get('observaciones', '').strip()

            with conexion.cursor() as cursor:
                cursor.execute("""
                    UPDATE citas 
                    SET cliente_nombre = %s, fecha = %s, hora = %s, servicio_id = %s, mascota_nombre = %s, observaciones = %s
                    WHERE id = %s AND tenant_id = %s
                """, (nombre, fecha, hora, servicio_id, mascota_nombre, observaciones, cita_id, tenant_id))
            conexion.commit()
            flash('Cita actualizada correctamente.', 'success')
        except Exception as e:
            conexion.rollback()
            logger.error(f"Error editando cita ID={cita_id}: {e}")
            flash(f'Error al editar cita: {e}', 'error')
        finally:
            conexion.close()

    return redirect(url_for('citas'))


@app.route('/cita/<int:cita_id>/eliminar', methods=['POST', 'DELETE'])
@app.route('/citas/eliminar/<int:cita_id>', methods=['POST', 'DELETE'])
def eliminar_cita(cita_id):
    if session.get('rol') not in ['admin', 'dueño']:
        return jsonify({'success': False, 'error': 'Sin permisos'}), 403

    tenant_id = session.get('tenant_id', 1)
    conexion = obtener_conexion()
    try:
        with conexion.cursor() as cursor:
            cursor.execute("DELETE FROM citas WHERE id = %s AND tenant_id = %s", (cita_id, tenant_id))
        conexion.commit()
        return jsonify({'success': True, 'message': 'Cita eliminada permanentemente.'})
    except Exception as e:
        conexion.rollback()
        logger.error(f"Error eliminando cita ID={cita_id}: {e}")
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conexion.close()


# ─── COMPRAS Y PROVEEDORES ───────────────────────────────────────────────────

@app.route('/compras')
def compras():
    if 'rol' not in session or session['rol'] not in ['admin', 'dueño']:
        flash('Acceso denegado.', 'error')
        return redirect(url_for('dashboard'))
    
    lista = compras_controlador.obtener_compras() if hasattr(compras_controlador, 'obtener_compras') else []
    return render_template('compras.html', compras=lista)


# ─── CARRITO Y PAGO E-COMMERCE ───────────────────────────────────────────────

def _get_cart():
    return session.get('cart', {})

def _save_cart(cart):
    session['cart'] = cart
    session.modified = True

@app.route('/carrito')
def ver_carrito():
    cart = _get_cart()
    items = []
    total_qty = 0
    total_price = 0.0
    
    for id_str, data in list(cart.items()):
        qty = int(data.get('qty', 0))
        if id_str.startswith('service_'):
            try:
                servicio_id = int(id_str.replace('service_', ''))
                servicio = servicios_controlador.obtener_servicio_por_id(servicio_id) if hasattr(servicios_controlador, 'obtener_servicio_por_id') else None
                if not servicio:
                    cart.pop(id_str, None)
                    continue
                
                precio_unitario = float(servicio[3] if len(servicio) > 3 and servicio[3] else 0.0)
                subtotal = precio_unitario * qty
                
                items.append({
                    'id': f'service_{servicio_id}',
                    'nombre': f"🏥 {servicio[1]} (Servicio)",
                    'descripcion': servicio[2] if len(servicio) > 2 else 'Servicio veterinario',
                    'precio_unitario': precio_unitario, 'qty': qty, 'subtotal': subtotal, 'type': 'service', 'imagen': None
                })
                total_qty += qty
                total_price += subtotal
            except Exception as e:
                logger.warning(f"Error procesando servicio {id_str}: {e}")
                continue
        else:
            try:
                pid = int(id_str)
                producto = productos_controlador.obtener_producto_por_id(pid) if hasattr(productos_controlador, 'obtener_producto_por_id') else None
                if not producto:
                    cart.pop(id_str, None)
                    continue
                
                stock = int(producto[3] or 0)
                if qty > stock:
                    qty = stock
                    cart[id_str]['qty'] = qty
                
                precio_unitario = float(producto[4] if len(producto) > 4 and producto[4] else 0.0)
                subtotal = precio_unitario * qty
                
                items.append({
                    'id': pid, 'nombre': producto[1], 'imagen': producto[7] if len(producto) > 7 else None,
                    'precio_unitario': precio_unitario, 'qty': qty, 'subtotal': subtotal, 'stock': stock, 'type': 'product'
                })
                total_qty += qty
                total_price += subtotal
            except Exception as e:
                logger.warning(f"Error procesando producto {id_str}: {e}")
                continue
    
    _save_cart(cart)
    return render_template('carrito.html', items=items, total_qty=total_qty, total_price=total_price)


@app.route('/carrito/agregar', methods=['POST'])
def agregar_al_carrito():
    try:
        producto_id = int(request.form.get('producto_id'))
        cantidad = int(request.form.get('cantidad', 1))
    except Exception:
        flash('Datos inválidos para agregar al carrito.', 'error')
        return redirect(request.referrer or url_for('index'))

    producto = productos_controlador.obtener_producto_por_id(producto_id) if hasattr(productos_controlador, 'obtener_producto_por_id') else None
    if not producto:
        flash('Producto no encontrado.', 'error')
        return redirect(url_for('index'))

    stock = int(producto[3] or 0)
    if cantidad <= 0 or stock < 1:
        flash('Cantidad inválida o producto agotado.', 'warning')
        return redirect(url_for('index'))

    cart = _get_cart()
    key = str(producto_id)
    existing = int(cart.get(key, {}).get('qty', 0))
    new_qty = min(existing + cantidad, stock)

    cart[key] = {'qty': new_qty}
    _save_cart(cart)

    is_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest' or request.is_json
    total_qty = sum(int(v.get('qty', 0)) for v in session.get('cart', {}).values())
    message = f'Producto "{producto[1]}" agregado al carrito.'
    
    if is_ajax:
        return jsonify({'success': True, 'message': message, 'cart_count': total_qty})

    flash(message, 'success')
    return redirect(url_for('ver_carrito'))


@app.route('/carrito/eliminar', methods=['POST'])
def carrito_eliminar():
    try:
        producto_id = int(request.form.get('producto_id'))
    except (ValueError, TypeError):
        return 'Error: ID de producto inválido', 400

    cart = _get_cart()
    key = str(producto_id)
    if key in cart:
        del cart[key]
        _save_cart(cart)
        return 'OK', 200
    return 'Error: Producto no encontrado', 404


@app.route('/carrito/actualizar', methods=['POST'])
def carrito_actualizar():
    try:
        producto_id = int(request.form.get('producto_id'))
        cantidad = int(request.form.get('cantidad', 1))
    except (ValueError, TypeError):
        return 'Error: Datos inválidos', 400

    producto = productos_controlador.obtener_producto_por_id(producto_id) if hasattr(productos_controlador, 'obtener_producto_por_id') else None
    if not producto or cantidad <= 0:
        return 'Error al actualizar', 400

    cart = _get_cart()
    key = str(producto_id)
    stock = int(producto[3] or 0)
    cart[key] = {'qty': min(cantidad, stock)}
    _save_cart(cart)
    return 'OK', 200


@app.route('/checkout')
def checkout():
    cart = _get_cart()
    if not cart:
        flash('Tu carrito está vacío', 'warning')
        return redirect(url_for('ver_carrito'))
    
    items = []
    subtotal = 0.0
    for id_str, data in cart.items():
        try:
            pid = int(id_str)
        except Exception:
            continue
        producto = productos_controlador.obtener_producto_por_id(pid) if hasattr(productos_controlador, 'obtener_producto_por_id') else None
        if not producto:
            continue
        qty = int(data.get('qty', 0))
        precio_unitario = float(producto[4] if len(producto) > 4 and producto[4] else 0.0)
        subtotal_item = precio_unitario * qty
        items.append({
            'id': pid, 'nombre': producto[1], 'imagen': producto[7] if len(producto) > 7 else None,
            'qty': qty, 'precio_unitario': precio_unitario, 'subtotal': subtotal_item
        })
        subtotal += subtotal_item
    
    igv = round(subtotal * 0.18, 2)
    total_final = subtotal + igv
    
    return render_template('checkout.html', items=items, total=subtotal, igv=igv, total_final=total_final)


@app.route('/ventas/procesar', methods=['POST'])
def procesar_venta_pos():
    if 'rol' not in session:
        return jsonify({'success': False, 'error': 'Sesión no válida'}), 401

    data = request.get_json() or {}
    tenant_id = session.get('tenant_id', 1)
    
    # 1. Obtener datos del cliente enviados desde el POS
    raw_cliente_id = data.get('cliente_id') or data.get('id_cliente')
    cliente_nombre_input = (data.get('cliente_nombre') or data.get('nombre_cliente') or '').strip()

    final_cliente_id = None
    final_cliente_nombre = 'Cliente General'
    final_cliente_doc = None

    conexion = obtener_conexion()
    try:
        with conexion.cursor() as cursor:
            # 2. Si se envió un ID, verificar que el cliente exista en la base de datos
            if raw_cliente_id and str(raw_cliente_id).isdigit():
                cursor.execute("""
                    SELECT id, nombre, documento 
                    FROM clientes 
                    WHERE id = %s AND (tenant_id = %s OR tenant_id IS NULL) AND activo = TRUE
                """, (int(raw_cliente_id), tenant_id))
                
                cli = cursor.fetchone()
                if cli:
                    final_cliente_id = cli[0]
                    final_cliente_nombre = cli[1]
                    final_cliente_doc = cli[2]

            # 3. Si no vino ID pero escribieron un nombre distinto a 'Cliente General', buscar por nombre exacto
            elif cliente_nombre_input and cliente_nombre_input.lower() != 'cliente general':
                cursor.execute("""
                    SELECT id, nombre, documento 
                    FROM clientes 
                    WHERE LOWER(nombre) = LOWER(%s) AND (tenant_id = %s OR tenant_id IS NULL) AND activo = TRUE
                    LIMIT 1
                """, (cliente_nombre_input, tenant_id))
                
                cli = cursor.fetchone()
                if cli:
                    final_cliente_id = cli[0]
                    final_cliente_nombre = cli[1]
                    final_cliente_doc = cli[2]
                else:
                    # Si el nombre no existe en la BD, se asigna como nombre del comprobante pero SIN crear fila en clientes
                    final_cliente_nombre = cliente_nombre_input

            # 4. Generar número de venta único
            num_venta = f"VNT-{int(time.time())}"
            fecha_actual = datetime.now(ZONA_HORARIA_PERU)
            productos_json = json.dumps(data.get('productos', []), ensure_ascii=False)

            # 5. Insertar la venta (final_cliente_id será NULL si es Cliente General)
            cursor.execute("""
                INSERT INTO ventas (
                    numero_venta, fecha_venta, vendedor_id, vendedor_nombre,
                    cliente_id, cliente_nombre, cliente_documento,
                    productos, subtotal, igv, total, metodo_pago,
                    monto_recibido, cambio_entregado, estado, tenant_id
                ) VALUES (
                    %s, %s, %s, %s,
                    %s, %s, %s,
                    %s, %s, %s, %s, %s,
                    %s, %s, 'completada'::estado_venta_enum, %s
                ) RETURNING id
            """, (
                num_venta,
                fecha_actual,
                session.get('usuario_id'),
                session.get('username', 'Cajero'),
                final_cliente_id,  # Si es None, PostgreSQL inserta NULL automáticamente
                final_cliente_nombre,
                final_cliente_doc,
                productos_json,
                data.get('subtotal', 0),
                data.get('igv', 0),
                data.get('total', 0),
                data.get('metodo_pago', 'efectivo'),
                data.get('monto_recibido', 0),
                data.get('cambio_entregado', 0),
                tenant_id
            ))

            venta_id = cursor.fetchone()[0]

            # 6. Descontar stock de los productos vendidos
            for prod in data.get('productos', []):
                if prod.get('tipo') != 'servicio' and prod.get('id'):
                    cursor.execute("""
                        UPDATE productos 
                        SET cantidad = GREATEST(cantidad - %s, 0) 
                        WHERE id = %s AND (tenant_id = %s OR tenant_id IS NULL)
                    """, (prod.get('cantidad', 1), prod.get('id'), tenant_id))

        conexion.commit()
        return jsonify({'success': True, 'venta_id': venta_id, 'numero_venta': num_venta})

    except Exception as e:
        conexion.rollback()
        logger.error(f"Error procesando venta POS: {e}")
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conexion.close()

# ─── MÓDULO DE MASCOTAS ──────────────────────────────────────────────────────

# ─── MÓDULO DE MASCOTAS ──────────────────────────────────────────────────────

@app.route('/mascotas', endpoint='mascota')
@app.route('/mascotas/lista', endpoint='mascotas')
@app.route('/mascotas/listar', endpoint='listar_mascotas')
def listar_mascotas():
    if 'rol' not in session or session['rol'] not in ['admin', 'empleado', 'dueño']:
        flash('Acceso denegado.', 'error')
        return redirect(url_for('dashboard'))

    tenant_id = session.get('tenant_id', 1)
    mascotas = []
    clientes = []

    conexion = obtener_conexion()
    try:
        with conexion.cursor() as cursor:
            # 1. Obtener lista de mascotas con los datos de su dueño (cliente)
            cursor.execute("""
                SELECT 
                    m.id, 
                    m.nombre AS mascota_nombre, 
                    m.especie, 
                    COALESCE(m.raza, '') AS raza, 
                    COALESCE(m.edad, 0) AS edad, 
                    COALESCE(m.peso, 0.0) AS peso, 
                    COALESCE(m.color, '') AS color, 
                    COALESCE(m.genero::text, '') AS genero, 
                    COALESCE(m.esterilizado, false) AS esterilizado, 
                    COALESCE(m.observaciones, '') AS observaciones,
                    m.cliente_id,
                    COALESCE(c.nombre, 'Sin Dueño') AS cliente_nombre,
                    COALESCE(c.documento, '') AS cliente_documento,
                    COALESCE(c.telefono, '') AS cliente_telefono
                FROM mascotas m
                LEFT JOIN clientes c ON m.cliente_id = c.id
                WHERE m.tenant_id = %s AND m.activo = TRUE
                ORDER BY m.id DESC
            """, (tenant_id,))
            
            rows = cursor.fetchall()
            for r in rows:
                mascotas.append({
                    'id': r[0],
                    'nombre': r[1],
                    'especie': r[2],
                    'raza': r[3],
                    'edad': r[4],
                    'peso': float(r[5]),
                    'color': r[6],
                    'genero': r[7],
                    'esterilizado': r[8],
                    'observaciones': r[9],
                    'cliente_id': r[10],
                    'cliente_nombre': r[11],
                    'cliente_documento': r[12],
                    'cliente_telefono': r[13]
                })

            # 2. Obtener lista de clientes activos para el selector/dropdown
            cursor.execute("""
                SELECT id, nombre, COALESCE(documento, '') 
                FROM clientes 
                WHERE tenant_id = %s AND activo = TRUE 
                ORDER BY nombre ASC
            """, (tenant_id,))
            clientes_rows = cursor.fetchall()
            for c in clientes_rows:
                clientes.append({
                    'id': c[0],
                    'nombre': c[1],
                    'documento': c[2]
                })

    except Exception as e:
        logger.error(f"Error al listar mascotas: {e}")
        flash(f"Error cargando mascotas: {str(e)}", "error")
    finally:
        conexion.close()

    return render_template('mascotas.html', mascotas=mascotas, clientes=clientes)

@app.route('/mascotas/crear', methods=['POST'])
def crear_mascota():
    if session.get('rol') not in ['admin', 'empleado', 'dueño']:
        if request.is_json:
            return jsonify({'success': False, 'error': 'Sin permisos'}), 403
        flash('Sin permisos.', 'error')
        return redirect(url_for('listar_mascotas'))

    tenant_id = session.get('tenant_id', 1)
    is_ajax = request.is_json or request.headers.get('X-Requested-With') == 'XMLHttpRequest'
    data = request.get_json(silent=True) or request.form

    # Captura de campos
    cliente_id = data.get('cliente_id')
    nombre = (data.get('nombre') or '').strip()
    especie = (data.get('especie') or '').strip()
    raza = (data.get('raza') or '').strip()
    color = (data.get('color') or '').strip()
    genero = (data.get('genero') or '').strip()
    observaciones = (data.get('observaciones') or '').strip()

    # Conversión de tipos numéricos y booleanos
    try:
        edad = int(data.get('edad')) if data.get('edad') not in [None, ''] else None
    except ValueError:
        edad = None

    try:
        peso = float(data.get('peso')) if data.get('peso') not in [None, ''] else None
    except ValueError:
        peso = None

    esterilizado = str(data.get('esterilizado')).lower() in ['true', '1', 'on', 'yes']

    # Validaciones obligatorias
    if not cliente_id or not nombre or not especie:
        msg = 'El cliente (dueño), nombre y especie son obligatorios.'
        if is_ajax:
            return jsonify({'success': False, 'error': msg}), 400
        flash(msg, 'error')
        return redirect(url_for('listar_mascotas'))

    conexion = obtener_conexion()
    try:
        with conexion.cursor() as cursor:
            # Validar que el cliente pertenezca al mismo tenant
            cursor.execute("SELECT id FROM clientes WHERE id = %s AND tenant_id = %s AND activo = TRUE", (cliente_id, tenant_id))
            if not cursor.fetchone():
                msg = 'El cliente seleccionado no existe o no es válido.'
                if is_ajax:
                    return jsonify({'success': False, 'error': msg}), 400
                flash(msg, 'error')
                return redirect(url_for('listar_mascotas'))

            # Insertar mascota
            cursor.execute("""
                INSERT INTO mascotas (
                    cliente_id, nombre, especie, raza, edad, peso, color, 
                    genero, esterilizado, observaciones, tenant_id, activo
                ) VALUES (
                    %s, %s, %s, %s, %s, %s, %s, 
                    NULLIF(%s, '')::genero_enum, %s, %s, %s, TRUE
                ) RETURNING id
            """, (
                cliente_id, nombre, especie, raza, edad, peso, color,
                genero, esterilizado, observaciones, tenant_id
            ))
            mascota_id = cursor.fetchone()[0]

        conexion.commit()
        msg = f'Mascota "{nombre}" registrada correctamente.'
        if is_ajax:
            return jsonify({'success': True, 'message': msg, 'mascota_id': mascota_id})
        flash(msg, 'success')

    except Exception as e:
        conexion.rollback()
        logger.error(f"Error al crear mascota: {e}")
        if is_ajax:
            return jsonify({'success': False, 'error': str(e)}), 500
        flash(f'Error al registrar mascota: {str(e)}', 'error')
    finally:
        conexion.close()

    return redirect(url_for('listar_mascotas'))


@app.route('/mascotas/editar/<int:mascota_id>', methods=['POST'])
def editar_mascota(mascota_id):
    if session.get('rol') not in ['admin', 'empleado', 'dueño']:
        flash('Sin permisos.', 'error')
        return redirect(url_for('listar_mascotas'))

    tenant_id = session.get('tenant_id', 1)
    data = request.form

    cliente_id = data.get('cliente_id')
    nombre = (data.get('nombre') or '').strip()
    especie = (data.get('especie') or '').strip()
    raza = (data.get('raza') or '').strip()
    color = (data.get('color') or '').strip()
    genero = (data.get('genero') or '').strip()
    observaciones = (data.get('observaciones') or '').strip()

    try:
        edad = int(data.get('edad')) if data.get('edad') not in [None, ''] else None
    except ValueError:
        edad = None

    try:
        peso = float(data.get('peso')) if data.get('peso') not in [None, ''] else None
    except ValueError:
        peso = None

    esterilizado = str(data.get('esterilizado')).lower() in ['true', '1', 'on', 'yes']

    if not cliente_id or not nombre or not especie:
        flash('El cliente (dueño), nombre y especie son obligatorios.', 'error')
        return redirect(url_for('listar_mascotas'))

    conexion = obtener_conexion()
    try:
        with conexion.cursor() as cursor:
            cursor.execute("""
                UPDATE mascotas 
                SET cliente_id = %s, nombre = %s, especie = %s, raza = %s, 
                    edad = %s, peso = %s, color = %s, genero = NULLIF(%s, '')::genero_enum, 
                    esterilizado = %s, observaciones = %s
                WHERE id = %s AND tenant_id = %s
            """, (
                cliente_id, nombre, especie, raza, edad, peso, color,
                genero, esterilizado, observaciones, mascota_id, tenant_id
            ))
        conexion.commit()
        flash('Información de la mascota actualizada.', 'success')
    except Exception as e:
        conexion.rollback()
        logger.error(f"Error al editar mascota ID={mascota_id}: {e}")
        flash(f'Error al editar mascota: {str(e)}', 'error')
    finally:
        conexion.close()

    return redirect(url_for('listar_mascotas'))


@app.route('/mascotas/eliminar/<int:mascota_id>', methods=['POST', 'DELETE'])
def eliminar_mascota(mascota_id):
    if session.get('rol') not in ['admin', 'dueño']:
        return jsonify({'success': False, 'error': 'Sin permisos'}), 403

    tenant_id = session.get('tenant_id', 1)
    conexion = obtener_conexion()
    try:
        with conexion.cursor() as cursor:
            # Borrado lógico desactivando la mascota
            cursor.execute("""
                UPDATE mascotas 
                SET activo = FALSE 
                WHERE id = %s AND tenant_id = %s
            """, (mascota_id, tenant_id))
        conexion.commit()
        return jsonify({'success': True, 'message': 'Mascota eliminada correctamente.'})
    except Exception as e:
        conexion.rollback()
        logger.error(f"Error al eliminar mascota ID={mascota_id}: {e}")
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conexion.close()


@app.route('/api/mascotas/buscar-clientes')
def buscar_clientes_para_mascota():
    if 'rol' not in session:
        return jsonify({'success': False, 'error': 'No autorizado'}), 401

    q = request.args.get('q', '').strip()
    tenant_id = session.get('tenant_id', 1)
    clientes = []

    conexion = obtener_conexion()
    try:
        with conexion.cursor() as cursor:
            if q:
                cursor.execute("""
                    SELECT id, nombre, COALESCE(documento, '')
                    FROM clientes
                    WHERE tenant_id = %s AND activo = TRUE
                      AND (LOWER(nombre) LIKE LOWER(%s) OR documento LIKE %s)
                    ORDER BY nombre ASC LIMIT 10
                """, (tenant_id, f'%{q}%', f'%{q}%'))
            else:
                cursor.execute("""
                    SELECT id, nombre, COALESCE(documento, '')
                    FROM clientes
                    WHERE tenant_id = %s AND activo = TRUE
                    ORDER BY nombre ASC LIMIT 10
                """, (tenant_id,))

            rows = cursor.fetchall()
            for r in rows:
                clientes.append({'id': r[0], 'nombre': r[1], 'documento': r[2]})

    except Exception as e:
        logger.error(f"Error buscando clientes para mascota: {e}")
    finally:
        conexion.close()

    return jsonify({'success': True, 'clientes': clientes})


# ─── FIDELIZACIÓN Y TENANTS ──────────────────────────────────────────────────

@app.route('/fidelizacion')
def fidelizacion():
    if 'rol' not in session or session['rol'] not in ['admin', 'empleado', 'dueño']:
        flash('Acceso denegado.', 'error')
        return redirect(url_for('dashboard'))
    
    alertas = fidelizacion_ctrl.obtener_alertas_recientes() if hasattr(fidelizacion_ctrl, 'obtener_alertas_recientes') else []
    return render_template('fidelizacion.html', alertas=alertas)


@app.route('/gestionar_tenants')
def gestionar_tenants():
    if session.get('rol') != 'superadmin':
        return redirect(url_for('login'))
    
    conexion = obtener_conexion()
    try:
        with conexion.cursor() as cursor:
            cursor.execute("SELECT id, nombre, slug, activo, created_at FROM tenants ORDER BY id")
            tenants = cursor.fetchall()
    except Exception as e:
        logger.error(f"Error listando tenants: {e}")
        tenants = []
    finally:
        conexion.close()
        
    return render_template('gestionar_tenants.html', tenants=tenants)


@app.route('/gestionar_tenants/crear', methods=['POST'])
def crear_tenant():
    if session.get('rol') != 'superadmin':
        return redirect(url_for('login'))
    
    nombre = request.form.get('nombre', '').strip()
    slug = request.form.get('slug', '').strip().lower().replace(' ', '_')
    admin_username = request.form.get('admin_username', '').strip()
    admin_password = request.form.get('admin_password', '').strip()
    
    if not all([nombre, slug, admin_username, admin_password]):
        flash('Todos los campos son obligatorios.', 'error')
        return redirect(url_for('gestionar_tenants'))
        
    from controladores.usuarios_controlador import hash_password
    conexion = obtener_conexion()
    try:
        with conexion.cursor() as cursor:
            cursor.execute("INSERT INTO tenants (nombre, slug) VALUES (%s, %s) RETURNING id", (nombre, slug))
            tenant_id = cursor.fetchone()[0]
            
            hashed = hash_password(admin_password)
            cursor.execute(
                "INSERT INTO usuarios (username, password, rol, activo, tenant_id) VALUES (%s, %s, 'admin', true, %s)",
                (admin_username, hashed, tenant_id)
            )
        conexion.commit()
        flash(f'Tenant "{nombre}" creado con éxito.', 'success')
    except Exception as e:
        conexion.rollback()
        logger.error(f"Error al crear tenant: {e}")
        flash(f'Error al crear tenant: {str(e)}', 'error')
    finally:
        conexion.close()

    return redirect(url_for('gestionar_tenants'))


@app.route('/gestionar_tenants/toggle/<int:tenant_id>', methods=['POST'])
def toggle_tenant(tenant_id):
    if session.get('rol') != 'superadmin':
        return redirect(url_for('login'))
        
    conexion = obtener_conexion()
    try:
        with conexion.cursor() as cursor:
            cursor.execute("UPDATE tenants SET activo = NOT activo WHERE id = %s", (tenant_id,))
        conexion.commit()
        flash('Estado del tenant actualizado.', 'success')
    except Exception as e:
        conexion.rollback()
        logger.error(f"Error cambiando estado tenant ID={tenant_id}: {e}")
        flash(f'Error: {str(e)}', 'error')
    finally:
        conexion.close()

    return redirect(url_for('gestionar_tenants'))


if __name__ == '__main__':
    app.run(debug=True)